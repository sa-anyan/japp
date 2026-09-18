from pathlib import Path
import re
import shutil
from openpyxl import load_workbook, Workbook
from datetime import date




# ============================================================
# 1. LOCATE DOWNLOADS
# ============================================================

downloads_folder = Path.home() / "Downloads"

# Public configuration: add filename tokens that identify your name.
# Example: ["jane", "doe"]
USER_NAME_TOKENS = ["your", "name"]


# ============================================================
# 2. FIND PDF AND DOCX FILES
# ============================================================

files = []

for file in downloads_folder.iterdir():

    if (
            file.is_file()
            and file.suffix.lower() in [".pdf", ".docx"]
            and not file.name.startswith("~$")
    ):
        files.append(file)


# ============================================================
# 3. CLASSIFY CVs AND COVER LETTERS
# ============================================================

cv_files = []
cover_letter_files = []
jd_files = []

for file in files:

    filename = file.stem.lower()

    if "cv" in filename or "resume" in filename:
        cv_files.append(file)

    elif "cover" in filename or "letter" in filename or "cl" in filename:
        cover_letter_files.append(file)

    elif "jd" in filename:
        jd_files.append(file)


# ============================================================
# 4. CHECK RESULTS
# ============================================================

print("\nCVs found:")
for file in cv_files:
    print(file.name)


print("\nCover letters found:")
for file in cover_letter_files:
    print(file.name)

print("\nJob descriptions found:")
for file in jd_files:
    print(file.name)

# ============================================================
# 5. PARSE COMPANY AND JOB TITLE
# ============================================================




def parse_filename(file):

    # Remove .docx / .pdf
    filename = file.stem

    # Remove duplicate markers added by Downloads
    # Example: "METLEN_EPC (2)" -> "METLEN_EPC"
    filename = re.sub(r"\s*\(\d+\)$", "", filename)

    # Treat "-" and "_" the same way
    filename = filename.replace("-", "_")

    # Split filename into pieces
    parts = filename.split("_")

    # Words we do NOT want when looking for the company
    ignore_words = {
        *USER_NAME_TOKENS,
        "cv",
        "resume",
        "cover",
        "coverletter",
        "letter",
        "cl",
        "jd"
    }

    useful_parts = []

    for part in parts:

        if part.lower() not in ignore_words:
            useful_parts.append(part)

    # First useful part = company
    if len(useful_parts) >= 1:
        company = useful_parts[0]
    else:
        company = None

    # Everything after company = job title
    if len(useful_parts) >= 2:
        job_title = " ".join(useful_parts[1:])
    else:
        job_title = ""

    return company, job_title

# ============================================================
# 6. TEST PARSER
# ============================================================

print("\n\nPARSED FILES")
print("=" * 70)

for file in cv_files + cover_letter_files + jd_files:

    company, job_title = parse_filename(file)

    print(f"\nFile:    {file.name}")
    print(f"Company: {company}")
    print(f"Job:     {job_title}")


# ============================================================
# 7. GROUP FILES BY COMPANY
# ============================================================

company_groups = {}


# ------------------------------------------------------------
# Add CVs
# ------------------------------------------------------------

for file in cv_files:

    company, job_title = parse_filename(file)

    if company not in company_groups:
        company_groups[company] = {
            "cv": [],
            "cover_letter": [],
            "jd": []
        }

    company_groups[company]["cv"].append(file)


# ------------------------------------------------------------
# Add Cover Letters
# ------------------------------------------------------------

for file in cover_letter_files:

    company, job_title = parse_filename(file)

    if company not in company_groups:
        company_groups[company] = {
            "cv": [],
            "cover_letter": [],
            "jd": []
        }

    company_groups[company]["cover_letter"].append(file)


# ------------------------------------------------------------
# Add Job Descriptions
# ------------------------------------------------------------

for file in jd_files:

    company, job_title = parse_filename(file)

    if company not in company_groups:
        company_groups[company] = {
            "cv": [],
            "cover_letter": [],
            "jd": []
        }

    company_groups[company]["jd"].append(file)


# ============================================================
# 8. CHECK COMPANY GROUPS
# ============================================================

print("\n\nCOMPANY GROUPS")
print("=" * 70)

for company, documents in company_groups.items():

    print(f"\n{company}")

    print("  CV:")
    for file in documents["cv"]:
        print(f"    - {file.name}")

    print("  Cover Letter:")
    for file in documents["cover_letter"]:
        print(f"    - {file.name}")

    print("  Job Description:")
    for file in documents["jd"]:
        print(f"    - {file.name}")

# ============================================================
# 9. JAPPS FOLDER
# ============================================================

japps_folder = Path.home() / "JAPPS"

japps_folder.mkdir(parents=True, exist_ok=True)

print(f"\nJAPPS folder: {japps_folder}")

# ============================================================
# 10. CREATE COMPANY FOLDERS
# ============================================================

for company in company_groups:

    company_folder = japps_folder / company

    company_folder.mkdir(exist_ok=True)

    print(f"Ready: {company_folder}")

# ============================================================
# NORMALISE JOB TITLE FOR MATCHING
# ============================================================

def normalise_job(job_title):

    return (
        job_title
        .lower()
        .replace(" ", "")
        .replace("_", "")
        .replace("-", "")
    )

# ============================================================
# 11. GROUP FILES BY COMPANY + JOB
# ============================================================

job_groups = {}


for company, documents in company_groups.items():

    # --------------------------------------------------------
    # Build a list of unique jobs from CV filenames
    # --------------------------------------------------------

    cv_jobs = []

    for file in documents["cv"]:

        _, job_title = parse_filename(file)

        if not job_title:
            continue

        existing_job = None

        for known_job in cv_jobs:

            if normalise_job(known_job) == normalise_job(job_title):
                existing_job = known_job
                break

        if existing_job is None:
            cv_jobs.append(job_title)


    # --------------------------------------------------------
    # Add CVs to their job groups
    # --------------------------------------------------------

    for file in documents["cv"]:

        _, job_title = parse_filename(file)

        if not job_title:
            job_title = "Unspecified"

        else:

            for known_job in cv_jobs:

                if normalise_job(known_job) == normalise_job(job_title):
                    job_title = known_job
                    break

        key = (company, job_title)

        if key not in job_groups:
            job_groups[key] = {
                "cv": [],
                "cover_letter": [],
                "jd":[]
            }

        job_groups[key]["cv"].append(file)


    # --------------------------------------------------------
    # Match cover letters to CV jobs
    # --------------------------------------------------------

    for file in documents["cover_letter"]:

        _, cover_job = parse_filename(file)

        matched_job = None

        if cover_job:

            for cv_job in cv_jobs:

                normal_cover = normalise_job(cover_job)
                normal_cv = normalise_job(cv_job)

                if (
                    normal_cover in normal_cv
                    or normal_cv in normal_cover
                ):
                    matched_job = cv_job
                    break

        elif len(cv_jobs) == 1:

            matched_job = cv_jobs[0]

        if matched_job is None:
            matched_job = "Unspecified"

        key = (company, matched_job)

        if key not in job_groups:
            job_groups[key] = {
                "cv": [],
                "cover_letter": [],
                "jd":[]
            }

        job_groups[key]["cover_letter"].append(file)



    # --------------------------------------------------------
    # Match Job Descriptions to CV jobs
    # --------------------------------------------------------

    for file in documents["jd"]:

            _, jd_job = parse_filename(file)

            matched_job = None

            if jd_job:

                for cv_job in cv_jobs:

                    normal_jd = normalise_job(jd_job)
                    normal_cv = normalise_job(cv_job)

                    if (
                            normal_jd in normal_cv
                            or normal_cv in normal_jd
                    ):
                        matched_job = cv_job
                        break

            elif len(cv_jobs) == 1:

                matched_job = cv_jobs[0]

            if matched_job is None:
                matched_job = "Unspecified"

            key = (company, matched_job)

            if key not in job_groups:
                job_groups[key] = {
                    "cv": [],
                    "cover_letter": [],
                    "jd": []
                }

            job_groups[key]["jd"].append(file)


# ============================================================
# 12. PREVIEW JOB FOLDERS
# ============================================================

print("\n\nJOB GROUPS")
print("=" * 70)

for (company, job_title), documents in job_groups.items():

    print(f"\n{company} → {job_title}")

    for file in documents["cv"]:
        print(f"  CV: {file.name}")

    for file in documents["cover_letter"]:
        print(f"  CL: {file.name}")

    for file in documents["jd"]:
        print(f"  JD: {file.name}")

# ============================================================
# 13. CREATE JOB FOLDERS AND MOVE FILES
# ============================================================

print("\n\nMOVING FILES")
print("=" * 70)


for (company, job_title), documents in job_groups.items():

    job_folder = japps_folder / company / job_title

    job_folder.mkdir(parents=True, exist_ok=True)

    print(f"\n{company} → {job_title}")

    for file in documents["cv"]:

        destination = job_folder / file.name

        if destination.exists():
            print(f"  SKIPPED: {file.name}")

        else:
            shutil.move(str(file), str(destination))
            print(f"  MOVED CV: {file.name}")

    for file in documents["cover_letter"]:

        destination = job_folder / file.name

        if destination.exists():
            print(f"  SKIPPED: {file.name}")

        else:
            shutil.move(str(file), str(destination))
            print(f"  MOVED CL: {file.name}")

    for file in documents["jd"]:

        destination = job_folder / file.name

        if destination.exists():
            print(f"  SKIPPED: {file.name}")

        else:
            shutil.move(str(file), str(destination))
            print(f"  MOVED JD: {file.name}")

# ============================================================
# 14. SCAN ENTIRE JAPPS FOLDER
# ============================================================

applications = []


for company_folder in japps_folder.iterdir():

    if not company_folder.is_dir():
        continue

    company = company_folder.name

    direct_cv = False
    direct_cover_letter = False
    direct_jd = False

    job_folders = []

    for item in company_folder.iterdir():

        if item.is_dir():
            job_folders.append(item)
            continue

        if not item.is_file():
            continue

        filename = item.stem.lower()

        if "cv" in filename or "resume" in filename:
            direct_cv = True

        elif (
            "cover" in filename
            or "letter" in filename
            or "cl" in filename
        ):
            direct_cover_letter = True

        elif (
            "jd" in filename
            or "job description" in filename
            or "jobdescription" in filename
        ):
            direct_jd = True

    if direct_cv or direct_cover_letter or direct_jd:

        applications.append({
            "company": company,
            "job_title": "",
            "cv": direct_cv,
            "cover_letter": direct_cover_letter,
            "jd": direct_jd
        })

    for job_folder in job_folders:

        job_title = job_folder.name

        cv_found = False
        cover_letter_found = False
        jd_found = False

        for file in job_folder.rglob("*"):

            if not file.is_file():
                continue

            filename = file.stem.lower()

            if "cv" in filename or "resume" in filename:
                cv_found = True

            elif (
                "cover" in filename
                or "letter" in filename
                or "cl" in filename
            ):
                cover_letter_found = True

            elif (
                "jd" in filename
                or "job description" in filename
                or "jobdescription" in filename
            ):
                jd_found = True

        applications.append({
            "company": company,
            "job_title": job_title,
            "cv": cv_found,
            "cover_letter": cover_letter_found,
            "jd": jd_found
        })

    if (
        not job_folders
        and not direct_cv
        and not direct_cover_letter
        and not direct_jd
    ):

        applications.append({
            "company": company,
            "job_title": "",
            "cv": False,
            "cover_letter": False,
            "jd": False
        })

# ============================================================
# 15. OPEN OR CREATE EXCEL TRACKER
# ============================================================

tracker_file = japps_folder / "tracker.xlsx"


if tracker_file.exists():

    workbook = load_workbook(tracker_file)
    sheet = workbook.active

    print("\nTracker found.")

else:

    workbook = Workbook()
    sheet = workbook.active

    sheet.title = "Applications"

    sheet.append([
        "Company",
        "Job Title",
        "Date Applied",
        "CV",
        "Cover Letter",
        "Job Description",
        "Contacted",
        "Status",
        "Job Link",
        "Contact",
        "Notes"
    ])

    workbook.save(tracker_file)

    print("\nTracker created.")


# ============================================================
# 16. FIND APPLICATIONS ALREADY IN EXCEL
# ============================================================

existing_applications = set()

for row in sheet.iter_rows(min_row=2, values_only=True):

    company = row[0]
    job_title = row[1]

    if company:

        company_key = str(company).lower().strip()

        if job_title:
            job_key = normalise_job(str(job_title))
        else:
            job_key = ""

        existing_applications.add(
            (company_key, job_key)
        )


# ============================================================
# 17. ADD NEW APPLICATIONS FROM JAPPS
# ============================================================

today = date.today()

print("\n\nUPDATING EXCEL")
print("=" * 70)


for application in applications:

    company = application["company"]
    job_title = application["job_title"]

    company_key = company.lower().strip()
    job_key = normalise_job(job_title)

    application_key = (
        company_key,
        job_key
    )

    if application_key in existing_applications:

        print(f"ALREADY EXISTS: {company} → {job_title}")
        continue

    has_cv = "Yes" if application["cv"] else "No"

    has_cover_letter = (
        "Yes"
        if application["cover_letter"]
        else "No"
    )

    has_jd = "Yes" if application["jd"] else "No"

    sheet.append([
        company,
        job_title,
        today,
        has_cv,
        has_cover_letter,
        has_jd,
        "No",
        "Applied",
        "",
        "",
        ""
    ])

    existing_applications.add(application_key)

    print(f"ADDED: {company} → {job_title}")



# ============================================================
# 18. SAVE TRACKER
# ============================================================

workbook.save(tracker_file)

print("\nTracker updated successfully.")

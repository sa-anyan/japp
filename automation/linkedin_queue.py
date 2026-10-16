"""Refresh a manual LinkedIn copy/paste queue; never message LinkedIn."""
from __future__ import annotations

import json
import shutil
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .approvals import eligible_companies
from .state import StateStore
from .workbook import _atomic_save, _header_map, _find_row, ensure_pipeline_column

HEADERS = ['ID Number', 'Company', 'Contact', 'Role', 'LinkedIn Profile', 'Email Status',
           'Documents', 'CV File', 'Proposal File', 'Message to Copy', 'LinkedIn Status',
           'Sent Date', 'Notes']
MANUAL = {'Sent manually', 'Replied', 'Do not contact'}


def _read(path):
    return json.loads(path.read_text(encoding='utf-8')) if path.is_file() else {}


def _profile(value):
    value = str(value or '').strip()
    parsed = urlparse(value)
    if parsed.scheme == 'https' and (parsed.hostname or '').endswith('.linkedin.com') or (
        parsed.scheme == 'https' and parsed.hostname == 'linkedin.com'
    ):
        return value if parsed.path.startswith('/in/') else ''
    return ''


def refresh_linkedin_queue(config, state: StateStore):
    root = Path(config['paths']['private_repo_dir'])
    workbook = Path(config['paths']['workbook'])
    rows = eligible_companies(workbook, config['paths']['registry'])
    book = load_workbook(workbook)
    try:
        sheet = book['05 LinkedIn'] if '05 LinkedIn' in book.sheetnames else book.create_sheet('05 LinkedIn')
        if sheet.max_row > 1 and sheet.cell(1, 1).value != HEADERS[0]:
            raise ValueError('05 LinkedIn has an unexpected layout')
        for col, header in enumerate(HEADERS, 1):
            sheet.cell(1, col, header).font = Font(bold=True, color='FFFFFF')
            sheet.cell(1, col).fill = PatternFill('solid', fgColor='17365D')
        existing = {(str(sheet.cell(r, 1).value), str(sheet.cell(r, 3).value).casefold()): r
                    for r in range(2, sheet.max_row + 1)}
        execution = book['03 Execution']
        ensure_pipeline_column(execution)
        columns = _header_map(execution)
        status_col = columns.get('LinkedIn Status') or execution.max_column + 1
        execution.cell(1, status_col, 'LinkedIn Status')
        execution.column_dimensions[execution.cell(1, status_col).column_letter].width = 30
        updated = 0
        for company in rows:
            folders = list((root / 'companies').glob(company.company_id + '_*'))
            if len(folders) != 1:
                continue
            folder = folders[0]; output = folder / 'outputs'
            leads = _read(root / 'research/contact_candidates' / f'{company.company_id}.json')
            people = state.list_contacts(company.company_id)
            verified_profiles = {_profile(p.get('linkedin_url')).rstrip('/').casefold() for p in people
                                 if p.get('email') and p.get('email_status') == 'verified'} - {''}
            people += [dict(x, full_name=x.get('full_name', '')) for x in
                       leads.get('candidates', []) + leads.get('unverified_people', [])]
            messages = _read(output / 'linkedin_outreach.json').get('messages', [])
            seen = set(); statuses = []
            for person in people:
                name = person.get('full_name', '').strip(); key = (company.company_id, name.casefold())
                if not name or key in seen: continue
                seen.add(key)
                if person.get('email') and person.get('email_status') == 'verified': continue
                if _profile(person.get('linkedin_url')).rstrip('/').casefold() in verified_profiles: continue
                r = existing.get(key) or sheet.max_row + 1
                old = [sheet.cell(r, col).value for col in range(1, len(HEADERS)+1)]
                profile = _profile(person.get('linkedin_url'))
                message = next((x.get('body', '') for x in messages
                                if str(x.get('full_name', '')).casefold() == name.casefold()
                                and _profile(x.get('linkedin_url')) == profile and profile), '')
                cv, proposal = output / 'tailored-cv.docx', output / 'proposal.docx'
                named_cvs = list(output.glob('Samuel Anyan_*_CV.docx'))
                named_proposals = list(output.glob('Samuel Anyan_*_Proposal.docx'))
                if len(named_cvs) == 1: cv = named_cvs[0]
                if len(named_proposals) == 1: proposal = named_proposals[0]
                docs = cv.is_file() and proposal.is_file()
                ready = bool(profile and message and docs)
                status = ('Do not contact' if person.get('suppressed') else
                          old[10] if old[10] in MANUAL else
                          'Ready to copy' if ready else 'Profile needed' if not profile else
                          'Documents needed' if not docs else 'Message needed')
                statuses.append(status)
                note = ('Review draft documents. Copy the message, send manually, then change LinkedIn Status. '
                        'This does not mark an email as sent.')
                values = [company.company_id, company.company, name, person.get('role', ''), profile,
                          'No email available' if not person.get('email') else person.get('email_status', 'unknown'),
                          'Available — review drafts' if docs else 'Pending',
                          str(cv.resolve()) if cv.is_file() else '', str(proposal.resolve()) if proposal.is_file() else '',
                          old[9] if old[10] in MANUAL else message, status, old[11], old[12] or note]
                for col, value in enumerate(values, 1):
                    cell = sheet.cell(r, col); cell.value = value
                    if isinstance(value, str): cell.data_type = 's'
                    cell.alignment = Alignment(vertical='top', wrap_text=True)
                    if col in (5, 8, 9) and value:
                        cell.hyperlink = value; cell.font = Font(color='0563C1', underline='single')
                sheet.cell(r, 11).fill = PatternFill('solid', fgColor='D9EAD3' if status == 'Ready to copy' else 'FFF2CC')
                sheet.cell(r, 12).number_format = 'dd mmm yyyy'
                sheet.row_dimensions[r].height = 300
                updated += 1
            erow = _find_row(execution, columns['ID Number'], company.company_id)
            if erow and statuses:
                execution.cell(erow, status_col, 'Ready to copy' if 'Ready to copy' in statuses else statuses[0])
            elif erow and verified_profiles:
                execution.cell(erow, status_col, 'Email available')
            for rr in range(2, sheet.max_row + 1):
                if (str(sheet.cell(rr, 1).value) == company.company_id
                        and _profile(sheet.cell(rr, 5).value).rstrip('/').casefold() in verified_profiles
                        and sheet.cell(rr, 11).value not in MANUAL):
                    sheet.cell(rr, 6, 'Verified email available')
                    sheet.cell(rr, 10, '')
                    sheet.cell(rr, 11, 'Email available')
        sheet.freeze_panes = 'C2'; sheet.auto_filter.ref = f'A1:M{max(2, sheet.max_row)}'
        sheet.sheet_view.showGridLines = False
        widths = [12, 38, 22, 30, 42, 24, 28, 32, 32, 85, 24, 18, 55]
        for col, width in enumerate(widths, 1): sheet.column_dimensions[sheet.cell(1, col).column_letter].width = width
        sheet.data_validations.dataValidation = []
        validation = DataValidation(type='list', formula1='"Profile needed,Documents needed,Message needed,Ready to copy,Email available,Sent manually,Replied,Do not contact"', allow_blank=True, showErrorMessage=True)
        validation.add(f'K2:K{max(1201, sheet.max_row)}'); sheet.add_data_validation(validation)
        backup = Path(config['paths']['backup_dir']) / ('linkedin-' + datetime.now().strftime('%Y%m%d-%H%M%S-%f') + '.xlsx')
        backup.parent.mkdir(parents=True, exist_ok=True); shutil.copy2(workbook, backup)
        _atomic_save(book, workbook)
        return {'rows': updated, 'sheet': '05 LinkedIn', 'backup': str(backup)}
    finally:
        book.close()

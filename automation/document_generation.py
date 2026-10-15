"""Deterministic application packages and attended AI document handoff."""

from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path
from typing import Any

from docx import Document


PLACEHOLDER_MARKERS = (
    "<!-- add your prompt here -->",
    "add your own prompt here",
    "TODO: add your own prompt",
)
PROMPT_FIELDS = {"company", "role", "job_description", "research_summary"}


class PromptNotConfigured(ValueError):
    pass


class DocumentValidationError(ValueError):
    pass


def safe_component(value: str) -> str:
    value = re.sub(r"[<>:\"/\\|?*]", " ", value)
    value = re.sub(r"[^\w -]", "", value, flags=re.UNICODE)
    return re.sub(r"\s+", " ", value).strip(" .") or "Unknown"


def expected_filename(company: str, role: str, kind: str) -> str:
    if kind == "CoverLetter":
        kind = "Cover Letter"
    if kind not in {"CV", "Cover Letter"}:
        raise ValueError("kind must be CV or Cover Letter")
    return f"Samuel Anyan_{safe_component(company)}_{safe_component(role)}_{kind}.docx"


def _research_record(research_dir: Path, company_id: str) -> tuple[Path, dict[str, Any]]:
    for path in sorted(research_dir.glob("*.json")):
        try:
            record = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, UnicodeError, json.JSONDecodeError):
            continue
        if str(record.get("company_id", "")).strip() == str(company_id).strip():
            return path, record
    raise FileNotFoundError(f"No valid research JSON found for company ID {company_id}")


def _prompt(path: Path) -> str | None:
    if not path.is_file():
        return None
    content = path.read_text(encoding="utf-8").strip()
    if not content or any(marker in content.casefold() for marker in PLACEHOLDER_MARKERS):
        return None
    return content


def _render(template: str, values: dict[str, str]) -> str:
    def replace(match: re.Match[str]) -> str:
        field = match.group(1)
        return values[field] if field in values else match.group(0)

    rendered = re.sub(r"\{([a-z_]+)\}", replace, template)
    unresolved = re.findall(r"\{([a-z_]+)\}", rendered)
    if unresolved:
        raise PromptNotConfigured(f"Prompt contains unsupported or missing placeholders: {', '.join(sorted(unresolved))}")
    return rendered


def build_input_package(company_id: int | str, config: dict[str, Any], *, role: str = "",
                        job_description: str = "") -> dict[str, Any]:
    paths = config["paths"]
    research_path, research = _research_record(Path(paths["research_dir"]), str(company_id))
    company = str(research["company"])
    role = role.strip() or "Speculative Approach"
    product = research.get("proposed_product") or {}
    research_summary = "\n".join(
        f"{key}: {research.get(key, '')}" for key in (
            "business_model", "entry_level_fit", "current_workflow", "opportunity"
        ) if research.get(key)
    )
    prompts_dir = Path(paths["prompts_dir"])
    cv_template = _prompt(prompts_dir / "cv_tailoring_prompt.md")
    proposal_template = _prompt(prompts_dir / "proposal_prompt.md")
    values = {
        "company": company,
        "role": role,
        "job_description": job_description,
        "research_summary": research_summary,
    }
    missing_prompts = []
    rendered_prompts = {}
    for key, template in (("cv_tailoring", cv_template), ("proposal", proposal_template)):
        if template is None:
            missing_prompts.append(key)
        else:
            rendered_prompts[key] = _render(template, values)

    prompt_digest = hashlib.sha256(
        json.dumps(rendered_prompts, sort_keys=True, ensure_ascii=False).encode("utf-8")
    ).hexdigest()
    master_cv = Path(paths["master_cv"]) if paths.get("master_cv") else None
    missing_inputs = list(missing_prompts)
    if master_cv is None or not master_cv.is_file():
        missing_inputs.append("master_cv")
    sources = research.get("sources") or []
    package = {
        "schema_version": "1.0",
        "company_id": str(company_id),
        "company": company,
        "opportunity_type": "advertised_vacancy" if job_description else "speculative_approach",
        "role": role,
        "job_description": job_description,
        "research_summary": research_summary,
        "research": {
            "website": research.get("website"),
            "business_model": research.get("business_model"),
            "entry_level_fit": research.get("entry_level_fit"),
            "current_workflow": research.get("current_workflow"),
            "workflow_evidence": research.get("workflow_evidence"),
            "opportunity": research.get("opportunity"),
            "proposed_product": product,
            "impact_model": research.get("impact_model"),
            "confidence": research.get("confidence"),
            "existing_tools_check": research.get("existing_tools_check"),
        },
        "sources": sources,
        "master_cv_path": str(master_cv) if master_cv and master_cv.is_file() else None,
        "prompts": rendered_prompts,
        "prompt_version": prompt_digest,
        "missing_prompts": missing_prompts,
        "missing_inputs": missing_inputs,
        "status": "awaiting_user_inputs" if missing_inputs else "ready_for_attended_handoff",
        "research_file": research_path.name,
    }

    output_dir = Path(paths["generated_dir"]) / f"{company_id}_{safe_component(company)}"
    output_dir.mkdir(parents=True, exist_ok=True)
    package_path = output_dir / "input_package.json"
    package_path.write_text(json.dumps(package, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    package["package_path"] = str(package_path)
    package["output_dir"] = str(output_dir)
    return package


def validate_docx(path: Path | str, company: str, role: str = "") -> dict[str, Any]:
    path = Path(path)
    if path.suffix.casefold() != ".docx" or not path.is_file() or path.stat().st_size == 0:
        raise DocumentValidationError(f"Document is missing, empty, or not DOCX: {path}")
    try:
        document = Document(path)
        text = "\n".join(
            [paragraph.text for paragraph in document.paragraphs]
            + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        )
    except Exception as exc:
        raise DocumentValidationError(f"Cannot open DOCX {path}: {exc}") from exc
    if company.casefold() not in text.casefold():
        raise DocumentValidationError(f"Document does not mention company {company!r}: {path.name}")
    if role and role.casefold() not in text.casefold():
        raise DocumentValidationError(f"Document does not mention role {role!r}: {path.name}")
    return {"path": str(path), "bytes": path.stat().st_size, "company": company, "role": role}


def validate_document_pair(cv_path: Path | str, cover_letter_path: Path | str,
                           company: str, role: str = "") -> list[dict[str, Any]]:
    return [validate_docx(cv_path, company, role), validate_docx(cover_letter_path, company, role)]


def save_text_as_docx(text: str, path: Path | str) -> Path:
    """Save user/provider-returned text verbatim; do not synthesize document content."""
    if not text.strip():
        raise DocumentValidationError("Cannot create a DOCX from empty text")
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    for paragraph in text.splitlines():
        document.add_paragraph(paragraph)
    document.save(path)
    return path
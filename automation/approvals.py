"""Read user-controlled approvals from the workbook without modifying it."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .sync import _load_company_registry, normalize_id


@dataclass(frozen=True)
class ApprovedCompany:
    company_id: str
    company: str
    execution_status: str
    cv_tailoring: str


def _normalise_choice(value: Any) -> str:
    return str(value or "").strip().casefold()


def read_execution_rows(workbook_path: Path | str, registry_path: Path | str | None = None
                        ) -> list[ApprovedCompany]:
    workbook_path = Path(workbook_path).expanduser()
    registry_path = (
        Path(registry_path).expanduser()
        if registry_path
        else workbook_path.with_name("JAPP_company_registry.csv")
    )
    registry = _load_company_registry(registry_path)
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook["03 Execution"]
        headers = {cell.value: cell.column - 1 for cell in sheet[1] if cell.value}
        required = {"ID Number", "Company", "Status", "CV Tailoring"}
        if not required.issubset(headers):
            raise ValueError(f"03 Execution is missing columns: {', '.join(sorted(required - headers.keys()))}")
        results = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            def value(header: str):
                column = headers[header]
                return row[column] if column < len(row) else None

            company_id = normalize_id(value("ID Number"))
            if company_id is None:
                continue
            company_value = value("Company")
            company = str(company_value).strip() if company_value not in (None, "") else registry.get(company_id, "")
            results.append(ApprovedCompany(
                company_id=company_id,
                company=company,
                execution_status=str(value("Status") or "").strip(),
                cv_tailoring=str(value("CV Tailoring") or "").strip(),
            ))
        return results
    finally:
        workbook.close()


def eligible_companies(workbook_path: Path | str, registry_path: Path | str | None = None
                       ) -> list[ApprovedCompany]:
    return [
        row for row in read_execution_rows(workbook_path, registry_path)
        if _normalise_choice(row.execution_status) == "approved"
        and _normalise_choice(row.cv_tailoring) == "start"
    ]
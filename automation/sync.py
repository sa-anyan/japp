"""Synchronize reviewed company research into the JAPP Excel dashboard."""

from __future__ import annotations

import hashlib
import json
import os
import shutil
import tempfile
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter


RESEARCH_FIELDS = {
    "business_model": "How They Make Money",
    "entry_level_fit": "How I Could Fit",
    "current_workflow": "Current Workflow",
}
RESEARCH_SHEET = "02 Research"
EXECUTION_SHEET = "03 Execution"
COMPANIES_SHEET = "01 Companies"


def normalize_id(value: Any) -> str | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value)
    if isinstance(value, str) and value.strip():
        return value.strip()
    return None


def value_hash(value: Any) -> str:
    payload = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str)
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def _text(value: Any, field: str, *, required: bool = False) -> str:
    if value is None and not required:
        return ""
    if not isinstance(value, str):
        raise ValueError(f"{field} must be a string")
    result = value.strip()
    if required and not result:
        raise ValueError(f"{field} is required")
    return result


def _parse_record(path: Path) -> dict[str, Any]:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise ValueError(f"cannot read JSON: {exc}") from exc
    if not isinstance(data, dict):
        raise ValueError("top-level JSON value must be an object")
    if str(data.get("schema_version", "")) != "1.0":
        raise ValueError("schema_version must be 1.0")
    company_id = normalize_id(data.get("company_id"))
    if company_id is None:
        raise ValueError("company_id is required and must be a string or number")
    company = _text(data.get("company"), "company", required=True)

    fields = {}
    for field in (*RESEARCH_FIELDS, "opportunity"):
        fields[field] = _text(data.get(field), field)

    product = data.get("proposed_product", {})
    if product is None:
        product = {}
    if not isinstance(product, dict):
        raise ValueError("proposed_product must be an object")
    product_name = _text(product.get("name"), "proposed_product.name")
    product_description = _text(product.get("description"), "proposed_product.description")
    product_summary = "\n".join(part for part in (product_name, product_description) if part)
    opportunity_parts = [part for part in (fields["opportunity"], product_summary) if part]

    return {
        "path": path,
        "company_id": company_id,
        "company_value": data["company_id"],
        "company": company,
        "research": {
            "Company": company,
            **{column: fields[field] for field, column in RESEARCH_FIELDS.items()},
            "SaaS / AI Improvement Opportunity": "\n\n".join(opportunity_parts),
        },
        "execution": {"Company": company, "SaaS Idea": product_summary},
    }


def _load_records(research_dir: Path) -> tuple[list[dict[str, Any]], list[str]]:
    if not research_dir.exists():
        return [], [f"Research directory not found: {research_dir}"]
    if not research_dir.is_dir():
        raise ValueError(f"Research path is not a directory: {research_dir}")

    records = []
    messages = []
    for path in sorted(research_dir.glob("*.json")):
        try:
            records.append(_parse_record(path))
        except ValueError as exc:
            messages.append(f"Malformed {path.name}: {exc}")

    by_id: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for record in records:
        by_id[record["company_id"]].append(record)
    duplicates = {company_id for company_id, matches in by_id.items() if len(matches) > 1}
    if duplicates:
        kept = []
        for record in records:
            if record["company_id"] in duplicates:
                messages.append(
                    f"Malformed {record['path'].name}: duplicate company_id {record['company_id']}"
                )
            else:
                kept.append(record)
        records = kept
    return records, messages


def _columns(sheet) -> dict[str, int]:
    columns = {}
    for cell in sheet[1]:
        if isinstance(cell.value, str) and cell.value.strip():
            if cell.value in columns:
                raise ValueError(f"Duplicate header {cell.value!r} in {sheet.title}")
            columns[cell.value] = cell.column
    return columns


def _require_columns(sheet, names: set[str]) -> dict[str, int]:
    columns = _columns(sheet)
    missing = sorted(names - columns.keys())
    if missing:
        raise ValueError(f"{sheet.title} is missing columns: {', '.join(missing)}")
    return columns


def _cell_value(sheet, cached_sheet, row: int, column: int) -> Any:
    cached = cached_sheet.cell(row, column).value
    value = sheet.cell(row, column).value
    if cached is not None:
        return cached
    if isinstance(value, str) and value.startswith("="):
        return None
    return value


def _id_index(sheet, cached_sheet, id_column: int) -> dict[str, int]:
    result = {}
    for row in range(2, max(sheet.max_row, cached_sheet.max_row) + 1):
        company_id = normalize_id(_cell_value(sheet, cached_sheet, row, id_column))
        if company_id is None:
            continue
        if company_id in result:
            raise ValueError(f"Duplicate company_id {company_id} in {sheet.title}")
        result[company_id] = row
    return result


def _load_manifest(path: Path) -> dict[str, str]:
    if not path.exists():
        return {}
    try:
        document = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise ValueError(f"Cannot read sync manifest {path}: {exc}") from exc
    cells = document.get("cells") if isinstance(document, dict) else None
    if not isinstance(cells, dict) or any(not isinstance(value, str) for value in cells.values()):
        raise ValueError(f"Invalid sync manifest format: {path}")
    return cells


def _is_formula(cell: Cell) -> bool:
    return cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("="))


def _write_cell(sheet, row: int, column: int, value: Any, manifest: dict[str, str]) -> bool:
    if value is None or value == "":
        return False
    cell = sheet.cell(row, column)
    if _is_formula(cell):
        return False
    key = f"{sheet.title}!{cell.coordinate}"
    current = cell.value
    owned = manifest.get(key) == value_hash(current)
    if current not in (None, "") and not owned:
        return False
    if current == value:
        return False
    cell.value = value
    manifest[key] = value_hash(value)
    return True


def _append_row(sheet, template_row: int | None, mapped_columns: set[int]) -> int:
    row = sheet.max_row + 1
    if template_row is None or template_row < 1:
        return row
    for column in range(1, sheet.max_column + 1):
        source = sheet.cell(template_row, column)
        target = sheet.cell(row, column)
        if source.has_style:
            target._style = source._style.copy()
        if column not in mapped_columns and _is_formula(source):
            target.value = Translator(source.value, origin=source.coordinate).translate_formula(
                target.coordinate
            )
    source_dimension = sheet.row_dimensions[template_row]
    target_dimension = sheet.row_dimensions[row]
    target_dimension.height = source_dimension.height
    target_dimension.hidden = source_dimension.hidden
    target_dimension.outlineLevel = source_dimension.outlineLevel
    target_dimension.collapsed = source_dimension.collapsed
    return row


def _extend_status_validation(sheet, row: int, status_column: int) -> None:
    reference = f"{get_column_letter(status_column)}{row}"
    for validation in sheet.data_validations.dataValidation:
        ranges = list(validation.sqref.ranges)
        applies_to_status = any(
            cell_range.min_col <= status_column <= cell_range.max_col for cell_range in ranges
        )
        already_covers_row = any(
            cell_range.min_col <= status_column <= cell_range.max_col
            and cell_range.min_row <= row <= cell_range.max_row
            for cell_range in ranges
        )
        if applies_to_status and not already_covers_row:
            validation.add(reference)


def _save_atomically(workbook, workbook_path: Path, manifest_path: Path,
                     manifest: dict[str, str], backup_dir: Path) -> Path:
    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S-%f")
    backup_dir.mkdir(parents=True, exist_ok=True)
    backup_path = backup_dir / f"{workbook_path.stem}-{timestamp}{workbook_path.suffix}"
    shutil.copy2(workbook_path, backup_path)

    workbook_temp = None
    manifest_temp = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=f".{workbook_path.stem}-", suffix=workbook_path.suffix,
            dir=workbook_path.parent, delete=False,
        ) as temp:
            workbook_temp = Path(temp.name)
        workbook.save(workbook_temp)

        manifest_path.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(
            mode="w", encoding="utf-8", prefix=f".{manifest_path.name}-",
            suffix=".tmp", dir=manifest_path.parent, delete=False,
        ) as temp:
            manifest_temp = Path(temp.name)
            json.dump({"version": 1, "cells": manifest}, temp, indent=2, sort_keys=True)
            temp.write("\n")

        os.replace(workbook_temp, workbook_path)
        workbook_temp = None
        os.replace(manifest_temp, manifest_path)
        manifest_temp = None
    except PermissionError as exc:
        raise RuntimeError(
            f"Cannot replace {workbook_path}; close the workbook in Excel and try again. "
            f"The backup is at {backup_path}."
        ) from exc
    except OSError as exc:
        raise RuntimeError(
            f"Could not safely save {workbook_path}: {exc}. The backup is at {backup_path}."
        ) from exc
    finally:
        for temporary in (workbook_temp, manifest_temp):
            if temporary is not None:
                temporary.unlink(missing_ok=True)
    return backup_path


def run_sync(research_dir: Path, workbook_path: Path, manifest_path: Path | None = None,
             backup_dir: Path = Path("backups"), dry_run: bool = False) -> dict[str, Any]:
    research_dir = Path(research_dir).expanduser()
    workbook_path = Path(workbook_path).expanduser()
    manifest_path = Path(manifest_path).expanduser() if manifest_path else workbook_path.with_suffix(".sync.json")
    backup_dir = Path(backup_dir).expanduser()
    if not workbook_path.is_file():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")

    records, messages = _load_records(research_dir)
    manifest = _load_manifest(manifest_path)
    workbook = load_workbook(workbook_path)
    cached_workbook = load_workbook(workbook_path, data_only=True)
    result = {
        "added": 0,
        "updated": 0,
        "skipped": 0,
        "malformed": sum(message.startswith("Malformed ") for message in messages),
        "messages": messages,
        "backup_path": None,
    }
    try:
        companies = workbook[COMPANIES_SHEET]
        research = workbook[RESEARCH_SHEET]
        execution = workbook[EXECUTION_SHEET]
        cached_companies = cached_workbook[COMPANIES_SHEET]
        cached_research = cached_workbook[RESEARCH_SHEET]
        cached_execution = cached_workbook[EXECUTION_SHEET]

        company_columns = _require_columns(companies, {"ID Number"})
        research_columns = _require_columns(
            research,
            {"ID Number", "Company", *RESEARCH_FIELDS.values(), "SaaS / AI Improvement Opportunity"},
        )
        execution_columns = _require_columns(
            execution, {"ID Number", "Company", "SaaS Idea", "Status", "CV Tailoring"}
        )
        company_ids = _id_index(companies, cached_companies, company_columns["ID Number"])
        research_ids = _id_index(research, cached_research, research_columns["ID Number"])
        execution_ids = _id_index(execution, cached_execution, execution_columns["ID Number"])

        research_mapped_columns = set(research_columns.values())
        execution_mapped_columns = {
            execution_columns[name] for name in ("ID Number", "Company", "SaaS Idea", "Status")
        }
        wrote_any = False

        for record in records:
            company_id = record["company_id"]
            if company_id not in company_ids:
                result["skipped"] += 1
                result["messages"].append(
                    f"Skipped {record['path'].name}: company_id {company_id} not found in 01 Companies"
                )
                continue

            research_row = research_ids.get(company_id)
            research_added = research_row is None
            if research_added:
                template = max(research_ids.values(), default=research.max_row or 1)
                research_row = _append_row(research, template, research_mapped_columns)
                research_ids[company_id] = research_row

            execution_row = execution_ids.get(company_id)
            execution_added = execution_row is None
            if execution_added:
                template = max(execution_ids.values(), default=execution.max_row or 1)
                execution_row = _append_row(execution, template, execution_mapped_columns)
                execution_ids[company_id] = execution_row
                _extend_status_validation(execution, execution_row, execution_columns["Status"])

            research_values = {
                "ID Number": record["company_value"],
                **record["research"],
            }
            execution_values = {
                "ID Number": record["company_value"],
                **record["execution"],
            }
            research_writes = 0
            execution_writes = 0
            for header, value in research_values.items():
                research_writes += _write_cell(
                    research, research_row, research_columns[header], value, manifest
                )
            for header, value in execution_values.items():
                execution_writes += _write_cell(
                    execution, execution_row, execution_columns[header], value, manifest
                )

            row_changed = research_added or execution_added or research_writes or execution_writes
            if row_changed:
                execution_writes += _write_cell(
                    execution, execution_row, execution_columns["Status"], "needs_review", manifest
                )

            changed = bool(research_added or execution_added or research_writes or execution_writes)
            wrote_any = wrote_any or changed
            if research_added or execution_added:
                result["added"] += 1
            elif changed:
                result["updated"] += 1
            else:
                result["skipped"] += 1

        if not records and not result["messages"]:
            result["messages"].append(f"No research JSON files found in {research_dir}")

        if wrote_any and not dry_run:
            result["backup_path"] = str(
                _save_atomically(workbook, workbook_path, manifest_path, manifest, backup_dir)
            )
    finally:
        workbook.close()
        cached_workbook.close()
    return result
"""Command-line entry point for the JAPP master automation engine."""

from __future__ import annotations

import argparse
import json
import sys
import webbrowser
from pathlib import Path
from typing import Any

from .approvals import eligible_companies
from .config import load_config
from .contacts import import_apollo_file
from .engine import request_stop, start_engine
from .orchestrator import monitor_handoffs, run_approved_queue, sync_research
from .outreach import (
    approve_email_batch,
    cancel_email_batch,
    create_email_batch,
    draft_email_batch,
    send_approved_batch,
)
from .providers.factory import create_mail_provider
from .state import StateStore
from .sync import run_sync


SITE_URLS = {
    "chatgpt": "https://chatgpt.com/",
    "apollo": "https://app.apollo.io/",
    "outlook": "https://outlook.office.com/mail/",
    "gmail": "https://mail.google.com/",
}


def _print_sync(result: dict[str, Any], dry_run: bool) -> None:
    for message in result["messages"]:
        print(message)
    print(
        f"{'Dry run ' if dry_run else ''}Summary: "
        f"{result['added']} added, {result['updated']} updated, "
        f"{result['skipped']} skipped, {result['malformed']} malformed"
    )
    if dry_run:
        print("Dry run complete; workbook, manifest, and backups were not written.")
    elif result.get("backup_path"):
        print(f"Workbook backup: {result['backup_path']}")


def _doctor(config: dict[str, Any]) -> int:
    issues = []
    if sys.version_info < (3, 11):
        issues.append("Python 3.11 or newer is required")
    workbook_path = Path(config["paths"]["workbook"])
    if not workbook_path.is_file():
        issues.append(f"Workbook not found: {workbook_path}")
    else:
        from openpyxl import load_workbook
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        required = {"01 Companies", "02 Research", "03 Execution", "04 Connections"}
        missing = required - set(workbook.sheetnames)
        workbook.close()
        if missing:
            issues.append(f"Workbook is missing sheets: {', '.join(sorted(missing))}")
    prompt_dir = Path(config["paths"]["prompts_dir"])
    for prompt in ("cv_tailoring_prompt.md", "proposal_prompt.md", "outreach_template.md"):
        if not (prompt_dir / prompt).is_file():
            issues.append(f"Prompt file is missing: {prompt_dir / prompt}")
    if config["mail"]["provider"] == "mock":
        print("Mail provider: mock; no account access or delivery.")
    if not config["mail"]["send_enabled"]:
        print("Live sending is disabled.")
    if not config["paths"].get("master_cv"):
        print("WARNING: master_cv is not configured; CV preparation will pause until set locally.")
    if issues:
        for issue in issues:
            print(f"ERROR: {issue}")
        return 1
    print("JAPP doctor: configuration and workbook checks passed.")
    return 0


def _status(config: dict[str, Any]) -> int:
    lock_path = Path(config["paths"]["lock_file"])
    if lock_path.exists():
        try:
            pid = lock_path.read_text(encoding="ascii").strip()
        except OSError:
            pid = "unknown"
        print(f"Engine lock present (PID {pid}).")
    else:
        print("Engine is stopped.")
    database = Path(config["paths"]["database"])
    if not database.is_file():
        print("State database is not initialized.")
        return 0
    state = StateStore(database)
    for company in state.list_companies():
        print(f"{company['company_id']} | {company['company_name']} | {company['state']}")
    services = state.services() if hasattr(state, "services") else []
    for service in services:
        print(f"service {service['service_name']}: {service['status']}")
    return 0


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="JAPP master automation engine")
    parser.add_argument("--config", type=Path, help="Optional local JSON config path")
    commands = parser.add_subparsers(dest="command", required=True)

    commands.add_parser("start", help="Start supervised local services")
    commands.add_parser("stop", help="Request graceful shutdown after the current step")
    commands.add_parser("status", help="Show engine and workflow status")
    commands.add_parser("doctor", help="Validate local configuration and workbook")
    exchange_parser = commands.add_parser("exchange", help="Private japp_data research and document workflow")
    exchange_commands = exchange_parser.add_subparsers(dest="exchange_command", required=True)
    exchange_commands.add_parser("research", help="Publish approved company research requests")
    exchange_commands.add_parser("package", help="Pull research and publish master CV, prompts and inputs")
    discovery = exchange_commands.add_parser("contacts", help="Find Apollo decision makers and reveal work emails")
    discovery.add_argument("--reveal-emails", action="store_true", required=True)
    discovery.add_argument("--max-reveals", type=int, required=True, help="Maximum new enrichment calls this run (not a credit-price guarantee)")
    exchange_commands.add_parser("collect", help="Pull, validate and organise completed output files")
    delivery = exchange_commands.add_parser("deliver", help="Collect results and create email drafts")
    delivery.add_argument("--send", action="store_true", help="Explicitly approve this batch and send; requires mail.send_enabled=true")
    dashboard = commands.add_parser("dashboard-copy", help="Add a Start selector to a new dashboard copy")
    dashboard.add_argument("--output", type=Path, required=True)
    dashboard.add_argument("--activate", action="store_true", help="Use the new copy in your local config")

    sync_parser = commands.add_parser("sync", help="Synchronize local research JSON into Excel")
    sync_parser.add_argument("--research-dir", type=Path)
    sync_parser.add_argument("--registry", type=Path)
    sync_parser.add_argument("--workbook", type=Path)
    sync_parser.add_argument("--manifest", type=Path)
    sync_parser.add_argument("--backup-dir", type=Path)
    sync_parser.add_argument("--dry-run", action="store_true")
    sync_parser.add_argument("--pull-github", action="store_true")

    run_parser = commands.add_parser("run", help="Process approved company work")
    run_parser.add_argument("--approved-only", action="store_true", default=True)
    run_parser.add_argument("--dry-run", action="store_true")
    run_parser.add_argument("--limit", type=int)
    run_parser.add_argument("--only")
    run_parser.add_argument("--batch-id")
    run_parser.add_argument("--send", action="store_true", help="Send only a previously approved batch")

    login_parser = commands.add_parser("login", help="Open a provider for manual sign-in")
    login_parser.add_argument("site", choices=tuple(SITE_URLS))

    contacts_parser = commands.add_parser("contacts", help="Manage local contact imports")
    contact_commands = contacts_parser.add_subparsers(dest="contacts_command", required=True)
    contact_import = contact_commands.add_parser("import-csv", help="Import an authorised Apollo CSV export")
    contact_import.add_argument("path", type=Path)
    for command_name, help_text in (("suppress", "Prevent any further email to a contact"),
                                    ("restore", "Remove a contact suppression")):
        contact_status = contact_commands.add_parser(command_name, help=help_text)
        contact_status.add_argument("contact_id", type=int)

    batch_parser = commands.add_parser("batch", help="Create and review outreach batches")
    batch_commands = batch_parser.add_subparsers(dest="batch_command", required=True)
    batch_commands.add_parser("create", help="Create an immutable batch preview")
    draft_parser = batch_commands.add_parser("draft", help="Create provider drafts from a batch")
    draft_parser.add_argument("batch_id")
    approve_parser = batch_commands.add_parser("approve", help="Approve the current batch snapshot")
    approve_parser.add_argument("batch_id")
    cancel_parser = batch_commands.add_parser("cancel", help="Delete unsent drafts and cancel a batch")
    cancel_parser.add_argument("batch_id")
    send_parser = batch_commands.add_parser("send", help="Send a previously approved batch")
    send_parser.add_argument("batch_id")
    send_parser.add_argument("--send", action="store_true", required=True)
    return parser


def main(argv: list[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    try:
        config = load_config(args.config)
        paths = config["paths"]

        if args.command == "dashboard-copy":
            from .dashboard import dashboard_copy
            output = dashboard_copy(Path(paths["workbook"]), args.output)
            if args.activate:
                config_path = Path(config["config_path"])
                local = json.loads(config_path.read_text(encoding="utf-8")) if config_path.is_file() else {}
                local.setdefault("paths", {})["workbook"] = str(output)
                config_path.parent.mkdir(parents=True, exist_ok=True)
                config_path.write_text(json.dumps(local, indent=2) + "\n", encoding="utf-8")
            print(output)
            return 0
        if args.command == "exchange":
            from .exchange import collect_outputs, export_research, package_inputs
            from .private_repository import PrivateRepository
            repo = PrivateRepository(config)
            repo.sync()
            if args.exchange_command == "research":
                result = export_research(config, repo)
            else:
                state = StateStore(paths["database"])
                state.initialize()
                if args.exchange_command == "package":
                    result = package_inputs(config, repo, state)
                elif args.exchange_command == "contacts":
                    from .contact_discovery import discover_contacts
                    result = discover_contacts(config, repo, state, reveal_emails=args.reveal_emails,
                                               max_reveals=args.max_reveals)
                elif args.exchange_command == "collect":
                    result = collect_outputs(config, repo, state)
                else:
                    from .delivery import deliver
                    result = deliver(config, repo, state, send=args.send)
            print(json.dumps(result, indent=2))
            return 0

        if args.command == "doctor":
            return _doctor(config)
        if args.command == "status":
            return _status(config)
        if args.command == "start":
            start_engine(config)
            return 0
        if args.command == "stop":
            if request_stop(paths["stop_file"], paths["lock_file"]):
                print("Stop requested; JAPP will exit after the current step.")
            else:
                print("JAPP is not running.")
            return 0
        if args.command == "sync":
            if args.pull_github and args.dry_run:
                parser.error("--pull-github cannot be combined with --dry-run because it updates local research files")
            overrides = any((args.research_dir, args.registry, args.workbook, args.manifest, args.backup_dir))
            if not overrides:
                if args.pull_github:
                    config["research"]["pull_before_sync"] = True
                result = sync_research(config, dry_run=args.dry_run)
                _print_sync(result, args.dry_run)
                return 0
            pull = None
            if args.pull_github or (config["research"]["pull_before_sync"] and not args.dry_run):
                from .github_research import pull_research_records
                pull = pull_research_records(
                    config["repository_root"],
                    remote=config["research"]["remote"],
                    branch=config["research"]["branch"],
                )
            result = run_sync(
                args.research_dir or paths["research_dir"],
                args.workbook or paths["workbook"],
                args.manifest or paths["manifest"],
                args.backup_dir or paths["backup_dir"],
                dry_run=args.dry_run,
                registry_path=args.registry or paths["registry"],
                source_commit=pull["commit"] if pull else None,
            )
            _print_sync(result, args.dry_run)
            return 0
        if args.command == "run":
            if args.send:
                if not args.batch_id:
                    parser.error("--send requires --batch-id")
                state = StateStore(paths["database"])
                state.initialize()
                provider = create_mail_provider(config)
                result = send_approved_batch(
                    config, state, args.batch_id, provider, send_flag=True
                )
                print(json.dumps(result, indent=2))
                return 0
            state = StateStore(paths["database"])
            if not args.dry_run:
                state.initialize()
                monitor_handoffs(config, state)
            result = run_approved_queue(
                config, state, dry_run=args.dry_run, limit=args.limit, only=args.only
            )
            print(json.dumps(result, indent=2))
            return 0
        if args.command == "login":
            webbrowser.open(SITE_URLS[args.site])
            print(f"Opened {args.site}. Sign in manually; JAPP does not read or store passwords.")
            return 0
        if args.command == "contacts":
            state = StateStore(paths["database"])
            state.initialize()
            if args.contacts_command in {"suppress", "restore"}:
                state.set_contact_suppressed(args.contact_id, args.contacts_command == "suppress")
                print(f"Contact {args.contact_id}: {args.contacts_command}d.")
                return 0
            result = import_apollo_file(
                args.path,
                registry_path=paths["registry"],
                workbook_path=paths["workbook"],
                state=state,
                target_titles=config["execution"]["target_titles"],
            )
            monitor_handoffs(config, state)
            print(json.dumps(result, indent=2))
            return 0
        if args.command == "batch":
            state = StateStore(paths["database"])
            state.initialize()
            if args.batch_command == "create":
                print(json.dumps(create_email_batch(config, state), indent=2))
                return 0
            if args.batch_command == "draft":
                result = draft_email_batch(config, state, args.batch_id, create_mail_provider(config))
                print(json.dumps(result, indent=2))
                return 0
            if args.batch_command == "approve":
                if not approve_email_batch(state, args.batch_id, paths["workbook"]):
                    parser.error("Batch is incomplete, changed, or not ready for approval")
                print(f"Batch {args.batch_id} approved.")
                return 0
            if args.batch_command == "cancel":
                cancel_email_batch(state, args.batch_id, create_mail_provider(config))
                print(f"Batch {args.batch_id} cancelled; unsent provider drafts were deleted.")
                return 0
            if args.batch_command == "send":
                result = send_approved_batch(
                    config, state, args.batch_id, create_mail_provider(config), send_flag=args.send
                )
                print(json.dumps(result, indent=2))
                return 0
    except (OSError, ValueError, RuntimeError, KeyError, TypeError) as exc:
        parser.exit(2, f"Error: {exc}\n")
    return 0

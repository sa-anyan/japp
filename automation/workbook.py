"""Safe, narrow workbook updates for pipeline status and contacts."""

from __future__ import annotations

import os
import re
import tempfile
from copy import copy
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .sync import normalize_id
from .pipeline_status import PIPELINE_OPTIONS, display_status, next_display_status


def _header_map(sheet) -> dict[str, int]:
    return {str(cell.value).strip(): cell.column for cell in sheet[1] if cell.value}


def _find_row(sheet, id_column: int, company_id: int | str) -> int | None:
    target = normalize_id(company_id)
    for row in range(2, sheet.max_row + 1):
        if normalize_id(sheet.cell(row, id_column).value) == target:
            return row
    return None


def _atomic_save(workbook, path: Path) -> None:
    temp_path = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=f".{path.stem}-", suffix=path.suffix, dir=path.parent, delete=False
        ) as temp:
            temp_path = Path(temp.name)
        workbook.save(temp_path)
        os.replace(temp_path, path)
        temp_path = None
    except PermissionError as exc:
        raise RuntimeError(f"Cannot update {path}; close it in Excel and try again") from exc
    finally:
        if temp_path:
            temp_path.unlink(missing_ok=True)


def ensure_pipeline_column(sheet) -> tuple[int, bool]:
    """Install the progress dropdown, migrate old labels and initialise ID rows."""
    columns = _header_map(sheet)
    changed = False
    column = columns.get("Pipeline Status")
    if column is None:
        column = sheet.max_column + 1
        cell = sheet.cell(1, column, "Pipeline Status")
        source = sheet.cell(1, columns["CV Tailoring"])
        if source.has_style:
            cell._style = copy(source._style)
        sheet.column_dimensions[get_column_letter(column)].width = 24
        changed = True
    letter = get_column_letter(column)
    if (sheet.column_dimensions[letter].width or 0) < 24:
        sheet.column_dimensions[letter].width = 24
        changed = True
    formula = '"' + ','.join(PIPELINE_OPTIONS) + '"'
    target = f"{letter}2:{letter}{max(sheet.max_row, 1201)}"
    matching = [v for v in sheet.data_validations.dataValidation
                if str(v.sqref) == target and v.formula1 == formula
                and v.showDropDown is False]
    if len(matching) != 1 or any(
        v not in matching and any(r.min_col <= column <= r.max_col for r in v.sqref.ranges)
        for v in sheet.data_validations.dataValidation
    ):
        for validation in list(sheet.data_validations.dataValidation):
            remaining = []
            for span in validation.sqref.ranges:
                if not span.min_col <= column <= span.max_col:
                    remaining.append(str(span))
                    continue
                if span.min_col < column:
                    remaining.append(f"{get_column_letter(span.min_col)}{span.min_row}:{get_column_letter(column-1)}{span.max_row}")
                if column < span.max_col:
                    remaining.append(f"{get_column_letter(column+1)}{span.min_row}:{get_column_letter(span.max_col)}{span.max_row}")
            if remaining:
                validation.sqref = ' '.join(remaining)
            else:
                sheet.data_validations.dataValidation.remove(validation)
        validation = DataValidation(type="list", formula1=formula, allow_blank=True,
                                    showDropDown=False, showErrorMessage=True)
        validation.error = "Choose a Pipeline Status from the list."
        validation.add(target)
        sheet.add_data_validation(validation)
        changed = True
    for row in range(2, sheet.max_row + 1):
        if normalize_id(sheet.cell(row, columns["ID Number"]).value) is None:
            continue
        cell = sheet.cell(row, column)
        value = display_status(cell.value)
        # Sent rows from older workbooks should also be protected before a rerun.
        if str(sheet.cell(row, columns["Status"]).value or "").strip().casefold() == "emailed":
            value = "Completed"
        if cell.value != value:
            cell.value = value
            changed = True
    return column, changed


def install_pipeline_status(workbook_path: Path | str) -> bool:
    path = Path(workbook_path).expanduser()
    workbook = load_workbook(path)
    try:
        _, changed = ensure_pipeline_column(workbook["03 Execution"])
        if changed:
            _atomic_save(workbook, path)
        return changed
    finally:
        workbook.close()


def update_pipeline_status(workbook_path: Path | str, company_id: int | str, status: str) -> bool:
    path = Path(workbook_path).expanduser()
    workbook = load_workbook(path)
    try:
        sheet = workbook["03 Execution"]
        columns = _header_map(sheet)
        column, changed = ensure_pipeline_column(sheet)
        row = _find_row(sheet, columns["ID Number"], company_id)
        if row is None:
            raise ValueError(f"Company ID {company_id} not found in 03 Execution")
        cell = sheet.cell(row, column)
        value = next_display_status(cell.value, status)
        if cell.value != value:
            cell.value = value
            changed = True
        if changed:
            _atomic_save(workbook, path)
        return changed
    finally:
        workbook.close()


def _slot_columns(columns: dict[str, int]) -> list[dict[str, int]]:
    slots = []
    for number in range(1, 6):
        slot = {}
        for field in ("Name", "Role", "Email", "LinkedIn", "Contact Status"):
            header = f"Person {number} {field}"
            if header not in columns:
                raise ValueError(f"04 Connections is missing column {header!r}")
            slot[field] = columns[header]
        slots.append(slot)
    return slots


def _contact_identity(name: Any, email: Any, linkedin: Any) -> tuple[str, str, str]:
    return (
        str(name or "").strip().casefold(),
        str(email or "").strip().casefold(),
        str(linkedin or "").strip().rstrip("/").casefold(),
    )


def write_contacts(workbook_path: Path | str, company_id: int | str,
                   contacts: Iterable[dict[str, Any]]) -> dict[str, int]:
    path = Path(workbook_path).expanduser()
    workbook = load_workbook(path)
    added = skipped = 0
    try:
        sheet = workbook["04 Connections"]
        columns = _header_map(sheet)
        if "ID Number" not in columns:
            raise ValueError("04 Connections is missing ID Number")
        row = _find_row(sheet, columns["ID Number"], company_id)
        if row is None:
            raise ValueError(f"Company ID {company_id} not found in 04 Connections")
        slots = _slot_columns(columns)

        for contact in contacts:
            name = str(contact.get("full_name") or "").strip()
            email = str(contact.get("email") or "").strip()
            linkedin = str(contact.get("linkedin_url") or "").strip()
            role = str(contact.get("role") or "").strip()
            if not name:
                skipped += 1
                continue
            candidate_identity = _contact_identity(name, email, linkedin)
            available = None
            for slot in slots:
                existing_name = sheet.cell(row, slot["Name"]).value
                existing_email = sheet.cell(row, slot["Email"]).value
                existing_linkedin = sheet.cell(row, slot["LinkedIn"]).value
                existing_identity = _contact_identity(existing_name, existing_email, existing_linkedin)
                if (
                    (candidate_identity[1] and candidate_identity[1] == existing_identity[1])
                    or (candidate_identity[2] and candidate_identity[2] == existing_identity[2])
                    or (candidate_identity[0] and candidate_identity[0] == existing_identity[0])
                ):
                    available = "duplicate"
                    break
                if not any((existing_name, existing_email, existing_linkedin)) and available is None:
                    available = slot
            if available == "duplicate" or available is None:
                skipped += 1
                continue
            status = str(contact.get("contact_status") or "").strip().casefold()
            if status not in {"found", "no_email", "unverified", "emailed"}:
                status = "found" if email else "no_email"
            sheet.cell(row, available["Name"]).value = name
            sheet.cell(row, available["Role"]).value = role
            sheet.cell(row, available["Email"]).value = email
            link_cell = sheet.cell(row, available["LinkedIn"])
            link_cell.value = linkedin
            if linkedin.startswith(("https://", "http://")):
                link_cell.hyperlink = linkedin
                link_cell.font = Font(name="Aptos", size=10, color="0563C1", underline="single")
            sheet.cell(row, available["Contact Status"]).value = status
            added += 1

        if added:
            _atomic_save(workbook, path)
        return {"added": added, "skipped": skipped}
    finally:
        workbook.close()


def mark_contact_emailed(workbook_path: Path | str, company_id: int | str, email: str) -> bool:
    path = Path(workbook_path).expanduser()
    workbook = load_workbook(path)
    changed = False
    try:
        sheet = workbook["04 Connections"]
        columns = _header_map(sheet)
        row = _find_row(sheet, columns["ID Number"], company_id)
        if row is None:
            raise ValueError(f"Company ID {company_id} not found in 04 Connections")
        for slot in _slot_columns(columns):
            if str(sheet.cell(row, slot["Email"]).value or "").strip().casefold() == email.strip().casefold():
                sheet.cell(row, slot["Contact Status"]).value = "emailed"
                changed = True
                break
        if changed:
            _atomic_save(workbook, path)
        return changed
    finally:
        workbook.close()


def contact_is_email_eligible(workbook_path: Path | str, company_id: int | str,
                              full_name: str, role: str, email: str) -> bool:
    workbook = load_workbook(Path(workbook_path).expanduser(), read_only=True, data_only=True)
    try:
        sheet = workbook["04 Connections"]
        columns = {str(cell.value).strip(): cell.column - 1 for cell in sheet[1] if cell.value}
        required = {"ID Number"} | {
            f"Person {number} {field}"
            for number in range(1, 6)
            for field in ("Name", "Role", "Email", "Contact Status")
        }
        if not required.issubset(columns):
            raise ValueError(f"04 Connections is missing columns: {', '.join(sorted(required - columns.keys()))}")
        target_id = normalize_id(company_id)
        for row in sheet.iter_rows(min_row=2, values_only=True):
            id_column = columns["ID Number"]
            if id_column >= len(row) or normalize_id(row[id_column]) != target_id:
                continue
            for number in range(1, 6):
                def value(field: str) -> str:
                    column = columns[f"Person {number} {field}"]
                    return str(row[column] or "").strip() if column < len(row) else ""
                if (
                    value("Name").casefold() == full_name.strip().casefold()
                    and value("Role").casefold() == role.strip().casefold()
                    and value("Email").casefold() == email.strip().casefold()
                ):
                    return value("Contact Status").casefold() not in {
                        "unverified", "no_email", "suppressed", "opted_out", "bounced", "emailed"
                    }
            return False
        return False
    finally:
        workbook.close()


def mark_execution_emailed(workbook_path: Path | str, company_id: int | str) -> bool:
    path = Path(workbook_path).expanduser()
    workbook = load_workbook(path)
    changed = False
    try:
        sheet = workbook["03 Execution"]
        columns = _header_map(sheet)
        row = _find_row(sheet, columns["ID Number"], company_id)
        if row is None:
            raise ValueError(f"Company ID {company_id} not found in 03 Execution")
        for header, value in (("Status", "emailed"), ("CV Tailoring", "done")):
            cell = sheet.cell(row, columns[header])
            if str(cell.value or "").strip().casefold() != value.casefold():
                cell.value = value
                changed = True
        column, installed = ensure_pipeline_column(sheet)
        changed = changed or installed
        if sheet.cell(row, column).value != "Completed":
            sheet.cell(row, column).value = "Completed"
            changed = True
        if changed:
            _atomic_save(workbook, path)
        return changed
    finally:
        workbook.close()

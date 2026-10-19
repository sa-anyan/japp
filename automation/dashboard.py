"""Install the Excel Start selector in a separate dashboard copy."""

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from .workbook import ensure_pipeline_column


def dashboard_copy(source: Path, destination: Path) -> Path:
    source, destination = source.expanduser().resolve(), destination.expanduser().resolve()
    if source == destination or destination.exists():
        raise ValueError("Choose a new output filename; dashboard setup never overwrites an existing workbook")
    book = load_workbook(source)
    try:
        sheet = book["03 Execution"]
        columns = {cell.value: cell.column for cell in sheet[1] if cell.value}
        column = columns["CV Tailoring"]
        letter = get_column_letter(column)
        # Remove only validations covering the targeted control cells. Keep any
        # unrelated ranges from the same validation object.
        for validation in list(sheet.data_validations.dataValidation):
            ranges = []
            for span in validation.sqref.ranges:
                if span.min_col <= column <= span.max_col:
                    if span.min_col < column:
                        ranges.append(f"{get_column_letter(span.min_col)}{span.min_row}:{get_column_letter(column-1)}{span.max_row}")
                    if column < span.max_col:
                        ranges.append(f"{get_column_letter(column+1)}{span.min_row}:{get_column_letter(span.max_col)}{span.max_row}")
                else:
                    ranges.append(str(span))
            if ranges:
                validation.sqref = " ".join(ranges)
            else:
                sheet.data_validations.dataValidation.remove(validation)
        selector = DataValidation(type="list", formula1='"Start,Done,Not required"',
                                  allow_blank=True, showDropDown=False, showErrorMessage=True)
        selector.promptTitle = "Start this company"
        selector.prompt = "Choose Start with Status Approved, then save and close Excel."
        selector.showInputMessage = True
        selector.error = "Choose Start, Done or Not required."
        selector.add(f"{letter}2:{letter}{max(sheet.max_row, 1201)}")
        sheet.add_data_validation(selector)
        sheet.cell(1, column).fill = PatternFill("solid", fgColor="D9EAD3")
        sheet.column_dimensions[letter].width = max(sheet.column_dimensions[letter].width or 12, 22)
        ensure_pipeline_column(sheet)
        destination.parent.mkdir(parents=True, exist_ok=True)
        book.save(destination)
    finally:
        book.close()
    return destination

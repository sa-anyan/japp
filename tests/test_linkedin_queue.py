import json
from openpyxl import load_workbook
from test_exchange import setup_exchange
from automation.exchange import company_folder
from automation.linkedin_queue import refresh_linkedin_queue


def test_linkedin_queue_preserves_manual_send_and_never_implies_email_sent(setup_exchange):
    config, repo, state = setup_exchange
    config['paths']['private_repo_dir'] = repo.path
    folder = company_folder(repo.path, '2', 'Example Energy') / 'outputs'
    folder.mkdir(parents=True)
    for name in ('tailored-cv.docx', 'proposal.docx'): (folder / name).write_bytes(b'fixture')
    profile = 'https://www.linkedin.com/in/a-person'
    state.save_contact('2', 'Example Energy', full_name='A Person', role='Director', email='', linkedin_url=profile, email_status='no_email')
    refresh_linkedin_queue(config, state)
    b = load_workbook(config['paths']['workbook']); assert b['05 LinkedIn']['K2'].value == 'Message needed'; b.close()
    (folder / 'linkedin_outreach.json').write_text(json.dumps({'messages': [{'full_name': 'A Person', 'linkedin_url': profile, 'body': 'Dear A, a proposed contribution.'}]}))
    refresh_linkedin_queue(config, state)
    b = load_workbook(config['paths']['workbook']); s = b['05 LinkedIn']
    assert s['K2'].value == 'Ready to copy'
    assert s['E2'].hyperlink.target == profile
    assert s['J2'].value == 'Dear A, a proposed contribution.'
    assert b['03 Execution']['F2'].value == 'Not Started'
    s['K2'] = 'Sent manually'; s['L2'] = '2026-10-16'; s['M2'] = 'My note'
    b.save(config['paths']['workbook']); b.close()
    refresh_linkedin_queue(config, state)
    b = load_workbook(config['paths']['workbook']); s = b['05 LinkedIn']
    assert s.max_row == 2
    assert [s.cell(2,c).value for c in (11,12,13)] == ['Sent manually','2026-10-16','My note']
    b.close()

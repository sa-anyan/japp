"""Readable dashboard progress and completion-based duplicate-send protection."""

from openpyxl import load_workbook

from test_automation import make_dashboard, prepare_outreach
from automation.approvals import eligible_companies
from automation.dashboard import dashboard_copy
from automation.outreach import approve_email_batch, create_email_batch, draft_email_batch, send_approved_batch
from automation.pipeline_status import PIPELINE_OPTIONS
from automation.providers.mail import MockMailProvider
from automation.workbook import install_pipeline_status, update_pipeline_status


def status(path):
    book = load_workbook(path)
    result = book["03 Execution"]["F2"].value
    book.close()
    return result


def test_installer_adds_dropdown_initialises_and_is_idempotent(tmp_path):
    path = tmp_path / "book.xlsx"
    make_dashboard(path)
    assert install_pipeline_status(path)
    assert status(path) == "Not Started"
    assert not install_pipeline_status(path)
    book = load_workbook(path)
    sheet = book["03 Execution"]
    assert sheet["F1"].value == "Pipeline Status"
    assert sheet["D2"].value == "Approved" and sheet["E2"].value == "Start"
    validation = [v for v in sheet.data_validations.dataValidation if "F2" in v]
    assert len(validation) == 1
    assert validation[0].formula1 == '"' + ','.join(PIPELINE_OPTIONS) + '"'
    assert len(validation[0].formula1) <= 255
    assert validation[0].showDropDown is False
    book.close()


def test_legacy_states_migrate_and_export_cannot_regress_progress(tmp_path):
    path = tmp_path / "book.xlsx"
    make_dashboard(path)
    book = load_workbook(path)
    book["03 Execution"]["F1"] = "Pipeline Status"
    book["03 Execution"]["F2"] = "AWAITING_AI_OUTPUT"
    book.save(path)
    book.close()
    install_pipeline_status(path)
    assert status(path) == "Researched"
    update_pipeline_status(path, "2", "Exported")
    assert status(path) == "Researched"
    update_pipeline_status(path, "2", "Contacts Ready")
    assert status(path) == "Contacts Ready"
    update_pipeline_status(path, "2", "DOCUMENTS_READY")
    assert status(path) == "Documents Ready"


def test_completed_row_is_excluded_even_if_approval_is_still_start(tmp_path):
    path = tmp_path / "book.xlsx"
    make_dashboard(path)
    update_pipeline_status(path, "2", "Completed")
    assert eligible_companies(path, tmp_path / "registry.csv") == []
    assert not update_pipeline_status(path, "2", "Exported")
    assert status(path) == "Completed"


def test_old_emailed_row_is_migrated_to_completed(tmp_path):
    path = tmp_path / "book.xlsx"
    make_dashboard(path, approvals=((2, "Example Energy", "emailed", "done"),))
    install_pipeline_status(path)
    assert status(path) == "Completed"


def test_completed_excel_marker_blocks_already_drafted_email(tmp_path):
    config, state = prepare_outreach(tmp_path)
    config["mail"]["send_enabled"] = True
    provider = MockMailProvider()
    batch = create_email_batch(config, state)
    draft_email_batch(config, state, batch["batch_id"], provider)
    assert status(config["paths"]["workbook"]) == "Email Drafted"
    assert approve_email_batch(state, batch["batch_id"], config["paths"]["workbook"])
    update_pipeline_status(config["paths"]["workbook"], "2", "Completed")
    result = send_approved_batch(config, state, batch["batch_id"], provider, send_flag=True)
    assert result == {"sent": 0, "skipped": 1, "failed": 0}
    assert provider.sent == []


def test_new_dashboard_copy_includes_pipeline_dropdown(tmp_path):
    source = tmp_path / "source.xlsx"
    make_dashboard(source)
    destination = dashboard_copy(source, tmp_path / "copy.xlsx")
    assert status(destination) == "Not Started"
    book = load_workbook(source)
    assert book["03 Execution"].max_column == 5
    book.close()


def test_sqlite_completion_blocks_new_batch_when_excel_label_is_reset(tmp_path):
    config, state = prepare_outreach(tmp_path)
    for value in ("RESEARCH_READY", "APPROVED", "APPLICATION_PACKAGE_READY", "AWAITING_AI_OUTPUT",
                  "DOCUMENTS_READY", "CONTACTS_READY", "EMAIL_DRAFTED", "AWAITING_BATCH_APPROVAL",
                  "EMAIL_QUEUED", "EMAIL_SUBMITTED", "COMPLETED"):
        state.transition("2", "Example Energy", value)
    update_pipeline_status(config["paths"]["workbook"], "2", "Not Started")
    batch = create_email_batch(config, state)
    assert batch["batch_id"] is None
    assert any("already completed" in reason for reason in batch["excluded"])

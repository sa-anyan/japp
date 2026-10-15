import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

from automation.sync import run_sync, value_hash


def make_workbook(path: Path, *, target_rows=False) -> None:
    workbook = Workbook()
    companies = workbook.active
    companies.title = "01 Companies"
    companies.append(["ID Number", "Company", "Town", "Sector"])
    companies.append([2, "Capital Energy Group Ltd", "London", "Energy"])

    research = workbook.create_sheet("02 Research")
    research.append([
        "ID Number", "Company", "How They Make Money", "How I Could Fit",
        "Current Workflow", "SaaS / AI Improvement Opportunity",
    ])
    execution = workbook.create_sheet("03 Execution")
    execution.append(["ID Number", "Company", "SaaS Idea", "Status", "CV Tailoring"])
    validation = DataValidation(
        type="list", formula1='"needs_review,approved,rejected"', allow_blank=True
    )
    validation.add("D2:D20")
    execution.add_data_validation(validation)
    execution.column_dimensions["C"].width = 42

    if target_rows:
        research.append([2, "Capital Energy Group Ltd", None, None, None, None])
        execution.append([2, "Capital Energy Group Ltd", None, None, None])
        research["A2"].font = Font(bold=True)

    workbook.save(path)


def write_research(directory: Path, *, company_id=2, company="Capital Energy Group Ltd",
                   business_model="Recurring energy services", product_name="Workflow Comparator",
                   product_description="Compares retrofit options.") -> Path:
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / f"{company_id}.json"
    path.write_text(json.dumps({
        "schema_version": "1.0",
        "company_id": company_id,
        "company": company,
        "research_status": "researched",
        "business_model": business_model,
        "entry_level_fit": "Commercial analysis",
        "current_workflow": "Spreadsheet review",
        "opportunity": "Automate comparison work.",
        "proposed_product": {"name": product_name, "description": product_description},
    }), encoding="utf-8")
    return path


def test_dry_run_does_not_change_workbook_or_create_sidecars(tmp_path):
    workbook_path = tmp_path / "dashboard.xlsx"
    research_dir = tmp_path / "research"
    make_workbook(workbook_path)
    write_research(research_dir)
    original = workbook_path.read_bytes()

    result = run_sync(research_dir, workbook_path, backup_dir=tmp_path / "backups", dry_run=True)

    assert result["added"] == 1
    assert result["updated"] == result["skipped"] == result["malformed"] == 0
    assert workbook_path.read_bytes() == original
    assert not workbook_path.with_suffix(".sync.json").exists()
    assert not (tmp_path / "backups").exists()


def test_adds_rows_preserves_validation_and_is_idempotent(tmp_path):
    workbook_path = tmp_path / "dashboard.xlsx"
    research_dir = tmp_path / "research"
    make_workbook(workbook_path)
    write_research(research_dir)

    first = run_sync(research_dir, workbook_path, backup_dir=tmp_path / "backups")
    assert first["added"] == 1
    workbook = load_workbook(workbook_path)
    research = workbook["02 Research"]
    execution = workbook["03 Execution"]
    assert research["C2"].value == "Recurring energy services"
    assert research["F2"].value == (
        "Automate comparison work.\n\n"
        "Workflow Comparator\nCompares retrofit options."
    )
    assert execution["C2"].value == "Workflow Comparator\nCompares retrofit options."
    assert execution["D2"].value == "needs_review"
    assert execution["E2"].value is None
    assert execution.column_dimensions["C"].width == 42
    assert any("D2" in str(validation.sqref) for validation in execution.data_validations.dataValidation)
    workbook.close()

    second = run_sync(research_dir, workbook_path, backup_dir=tmp_path / "backups")
    assert second["added"] == second["updated"] == 0
    assert second["skipped"] == 1
    assert len(list((tmp_path / "backups").glob("*.xlsx"))) == 1


def test_updates_owned_cells_but_preserves_manual_values_and_formulas(tmp_path):
    workbook_path = tmp_path / "dashboard.xlsx"
    research_dir = tmp_path / "research"
    make_workbook(workbook_path, target_rows=True)
    write_research(research_dir, business_model="Updated business model")
    workbook = load_workbook(workbook_path)
    research = workbook["02 Research"]
    execution = workbook["03 Execution"]
    research["C2"] = "Manual edit"
    research["D2"] = "=1+1"
    execution["D2"] = "approved"
    execution["E2"] = "start"
    workbook.save(workbook_path)

    manifest_path = workbook_path.with_suffix(".sync.json")
    manifest_path.write_text(json.dumps({
        "version": 1,
        "cells": {"02 Research!C2": value_hash("Old generated value")},
    }), encoding="utf-8")

    result = run_sync(research_dir, workbook_path, backup_dir=tmp_path / "backups")
    workbook = load_workbook(workbook_path, data_only=False)
    research = workbook["02 Research"]
    execution = workbook["03 Execution"]
    assert result["updated"] == 1
    assert research["C2"].value == "Manual edit"
    assert research["D2"].value == "=1+1"
    assert execution["D2"].value == "approved"
    assert execution["E2"].value == "start"
    assert research["E2"].value == "Spreadsheet review"
    workbook.close()


def test_reports_malformed_and_unknown_company_ids(tmp_path):
    workbook_path = tmp_path / "dashboard.xlsx"
    research_dir = tmp_path / "research"
    make_workbook(workbook_path)
    research_dir.mkdir()
    (research_dir / "broken.json").write_text("{", encoding="utf-8")
    write_research(research_dir, company_id=999, company="Unknown Company")

    result = run_sync(research_dir, workbook_path, dry_run=True)

    assert result["malformed"] == 1
    assert result["skipped"] == 1
    assert any("Malformed broken.json" in message for message in result["messages"])
    assert any("company_id 999 not found" in message for message in result["messages"])


def test_missing_research_directory_is_not_counted_as_malformed(tmp_path):
    workbook_path = tmp_path / "dashboard.xlsx"
    make_workbook(workbook_path)

    result = run_sync(tmp_path / "missing-research", workbook_path, dry_run=True)

    assert result["malformed"] == 0
    assert any("Research directory not found" in message for message in result["messages"])


def test_changed_owned_value_sets_needs_review_and_creates_backup(tmp_path):
    workbook_path = tmp_path / "dashboard.xlsx"
    research_dir = tmp_path / "research"
    make_workbook(workbook_path, target_rows=True)
    write_research(research_dir)
    manifest_path = workbook_path.with_suffix(".sync.json")
    manifest_path.write_text(json.dumps({
        "version": 1,
        "cells": {"02 Research!C2": value_hash("Old generated value")},
    }), encoding="utf-8")

    result = run_sync(research_dir, workbook_path, backup_dir=tmp_path / "backups")

    workbook = load_workbook(workbook_path)
    assert result["updated"] == 1
    assert workbook["02 Research"]["C2"].value == "Recurring energy services"
    assert workbook["03 Execution"]["D2"].value == "needs_review"
    assert Path(result["backup_path"]).is_file()
    workbook.close()
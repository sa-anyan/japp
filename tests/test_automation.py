import csv
import json
import tempfile
from pathlib import Path

import httpx
import pytest
from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

from automation.approvals import eligible_companies
from automation.config import load_config
from automation.contacts import Contact, import_apollo_csv, import_apollo_file, rank_contact
from automation.document_generation import (
    DocumentValidationError,
    build_input_package,
    expected_filename,
    save_text_as_docx,
    validate_docx,
)
from automation.engine import InstanceAlreadyRunning, SingleInstanceLock, request_stop, start_engine
from automation.github_research import DirtyResearchTree, pull_research_records
from automation.handoff import SupervisedChatGPTHandoff, classify_document, snapshot_downloads
from automation.orchestrator import import_pending_contacts, run_approved_queue
from automation.outreach import (
    BatchSafetyError,
    approve_email_batch,
    cancel_email_batch,
    create_email_batch,
    draft_email_batch,
    send_approved_batch,
)
from automation.providers.mail import GmailProvider, GraphOutlookProvider, MockMailProvider
from automation.state import InvalidTransition, StateStore
from automation.workbook import write_contacts


CONTACT_HEADERS = ["ID Number", "Company"] + [
    f"Person {number} {field}"
    for number in range(1, 6)
    for field in ("Name", "Role", "Email", "LinkedIn", "Contact Status")
]


def make_dashboard(path: Path, *, approvals=((2, "Example Energy", "Approved", "Start"),)) -> None:
    book = Workbook()
    companies = book.active
    companies.title = "01 Companies"
    companies.append(["ID Number", "Company", "Town", "Sector"])
    for company_id, company, _, _ in approvals:
        companies.append([company_id, company, "London", "Energy"])

    research = book.create_sheet("02 Research")
    research.append([
        "ID Number", "Company", "Business Idea", "How They Make Money",
        "How I Could Fit", "Current Workflow", "SaaS / AI Improvement Opportunity",
    ])
    execution = book.create_sheet("03 Execution")
    execution.append(["ID Number", "Company", "SaaS Idea", "Status", "CV Tailoring"])
    for company_id, company, status, tailoring in approvals:
        execution.append([company_id, company, None, status, tailoring])
    status_validation = DataValidation(
        type="list", formula1='"Pending Review,Approved,Rejected,emailed"', allow_blank=True
    )
    tailoring_validation = DataValidation(
        type="list", formula1='"Start,Done,Not required"', allow_blank=True
    )
    status_validation.add("D2:D1200")
    tailoring_validation.add("E2:E1200")
    execution.add_data_validation(status_validation)
    execution.add_data_validation(tailoring_validation)
    execution.column_dimensions["C"].width = 42

    connections = book.create_sheet("04 Connections")
    connections.append(CONTACT_HEADERS)
    for company_id, company, _, _ in approvals:
        connections.append([company_id, company] + [None] * 25)
    book.save(path)


def make_config(tmp_path: Path, workbook: Path) -> dict:
    config = load_config()
    paths = config["paths"]
    paths.update({
        "workbook": workbook,
        "research_dir": tmp_path / "research" / "companies",
        "registry": tmp_path / "registry.csv",
        "manifest": tmp_path / "state" / "sync.json",
        "database": tmp_path / "state" / "japp.sqlite",
        "backup_dir": tmp_path / "backups",
        "generated_dir": tmp_path / "generated",
        "contacts_dir": tmp_path / "contacts",
        "apollo_import_dir": tmp_path / "apollo_imports",
        "logs_dir": tmp_path / "logs",
        "prompts_dir": tmp_path / "prompts",
        "master_cv": None,
        "profile_dir": tmp_path / "profile",
        "stop_file": tmp_path / "STOP",
        "lock_file": tmp_path / ".japp.lock",
    })
    paths["research_dir"].mkdir(parents=True)
    paths["prompts_dir"].mkdir(parents=True)
    paths["registry"].write_text(
        "id,company,town,sector\n2,Example Energy,London,Energy\n",
        encoding="utf-8",
    )
    return config


def write_research(config: dict, company_id=2, company="Example Energy") -> Path:
    path = config["paths"]["research_dir"] / f"{company_id:04d}.json"
    path.write_text(json.dumps({
        "schema_version": "1.0",
        "company_id": company_id,
        "company": company,
        "research_status": "researched",
        "business_model": "Recurring contracts",
        "entry_level_fit": "Commercial analysis",
        "current_workflow": "Manual spreadsheet review",
        "opportunity": "Improve scenario comparisons.",
        "proposed_product": {
            "name": "Scenario Comparator",
            "description": "Compare options with transparent assumptions.",
        },
        "impact_model": {"metric": "validated time saved"},
        "confidence": "medium",
        "sources": [{"url": "https://example.test/source"}],
    }), encoding="utf-8")
    return path


def make_doc(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_paragraph(text)
    doc.save(path)


def prepare_outreach(tmp_path: Path, contact_count=1):
    workbook_path = tmp_path / "dashboard.xlsx"
    make_dashboard(workbook_path)
    config = make_config(tmp_path, workbook_path)
    output = config["paths"]["generated_dir"] / "2_Example Energy"
    output.mkdir(parents=True)
    make_doc(output / "Samuel Anyan_Example Energy_Speculative Approach_CV.docx", "Samuel Anyan analyst CV Education Skills")
    make_doc(output / "Samuel Anyan_Example Energy_Speculative Approach_Cover Letter.docx", "Example Energy introduction")
    (output / "introduction.md").write_text("I would be glad to discuss a practical scenario model.", encoding="utf-8")
    (output / "outreach.json").write_text(json.dumps({"messages": [
        {"recipient": f"person{index}@example.test", "subject": "Introduction: Example Energy",
         "body": f"Dear Person {index},\n\nI would be glad to discuss a practical scenario model.\n\nRegards, Samuel"}
        for index in range(contact_count)]}))
    state = StateStore(config["paths"]["database"])
    state.initialize()
    for index in range(contact_count):
        email = f"person{index}@example.test"
        contact_id, _ = state.save_contact(
            2, "Example Energy", full_name=f"Person {index}", role="Hiring Manager",
            email=email, email_status="verified", source="fixture",
        )
        slot = index + 1
        book = load_workbook(workbook_path)
        sheet = book["04 Connections"]
        sheet.cell(2, 2 + (slot - 1) * 5 + 1).value = f"Person {index}"
        sheet.cell(2, 2 + (slot - 1) * 5 + 2).value = "Hiring Manager"
        sheet.cell(2, 2 + (slot - 1) * 5 + 3).value = email
        sheet.cell(2, 2 + (slot - 1) * 5 + 5).value = "found"
        book.save(workbook_path)
        book.close()
    return config, state


def test_approval_requires_approved_and_start_and_is_read_only(tmp_path):
    workbook_path = tmp_path / "dashboard.xlsx"
    approvals = (
        (2, "Example Energy", "Approved", "Start"),
        (3, "Second Energy", "Approved", "Done"),
        (4, "Third Energy", "Pending Review", "Start"),
    )
    make_dashboard(workbook_path, approvals=approvals)
    registry = tmp_path / "registry.csv"
    registry.write_text("id,company,town,sector\n", encoding="utf-8")
    before = workbook_path.read_bytes()

    result = eligible_companies(workbook_path, registry)

    assert [(row.company_id, row.company) for row in result] == [("2", "Example Energy")]
    assert workbook_path.read_bytes() == before


def test_approval_reader_handles_sparse_trailing_cells(tmp_path):
    workbook_path = tmp_path / "sparse.xlsx"
    make_dashboard(workbook_path)
    book = load_workbook(workbook_path)
    sheet = book["03 Execution"]
    sheet["E2"] = None
    book.save(workbook_path)
    book.close()

    assert eligible_companies(workbook_path, tmp_path / "missing.csv") == []


def test_contact_writer_preserves_manual_slot_deduplicates_and_links(tmp_path):
    workbook_path = tmp_path / "dashboard.xlsx"
    make_dashboard(workbook_path)
    book = load_workbook(workbook_path)
    sheet = book["04 Connections"]
    sheet["C2"] = "Manual Contact"
    sheet["D2"] = "Director"
    sheet["E2"] = "manual@example.test"
    sheet["F2"] = "https://linkedin.com/in/manual"
    sheet["G2"] = "found"
    book.save(workbook_path)
    book.close()

    result = write_contacts(workbook_path, 2, [
        {"full_name": "Other Name", "role": "Hiring Manager", "email": "manual@example.test",
         "linkedin_url": "https://linkedin.com/in/other", "email_status": "verified"},
        {"full_name": "New Contact", "role": "Hiring Manager", "email": "new@example.test",
         "linkedin_url": "https://linkedin.com/in/new", "email_status": "verified"},
    ])

    book = load_workbook(workbook_path)
    sheet = book["04 Connections"]
    assert result == {"added": 1, "skipped": 1}
    assert sheet["C2"].value == "Manual Contact"
    assert sheet["E2"].value == "manual@example.test"
    assert sheet["H2"].value == "New Contact"
    assert sheet["J2"].value == "new@example.test"
    assert sheet["K2"].hyperlink.target == "https://linkedin.com/in/new"
    book.close()
    assert write_contacts(workbook_path, 2, [{"full_name": "New Contact", "email": "new@example.test"}]) == {
        "added": 0, "skipped": 1
    }


def test_state_machine_and_contact_idempotency(tmp_path):
    store = StateStore(tmp_path / "state.sqlite")
    store.initialize()
    store.transition(2, "Example Energy", "RESEARCH_READY")
    with pytest.raises(InvalidTransition):
        store.transition(2, "Example Energy", "COMPLETED")
    contact_id, created = store.save_contact(
        2, "Example Energy", full_name="A Person", email="a@example.test", email_status="verified"
    )
    same_id, created_again = store.save_contact(
        2, "Example Energy", full_name="A Person", email="A@example.test", email_status="verified"
    )
    assert created and not created_again and contact_id == same_id
    store.set_contact_suppressed(contact_id)
    assert store.list_contacts(2)[0]["suppressed"] == 1


def test_research_package_is_deterministic_and_prompts_remain_user_owned(tmp_path):
    workbook = tmp_path / "dashboard.xlsx"
    make_dashboard(workbook)
    config = make_config(tmp_path, workbook)
    write_research(config)
    cv_prompt = config["paths"]["prompts_dir"] / "cv_tailoring_prompt.md"
    cv_prompt.write_text("<!-- Add your own prompt here. -->", encoding="utf-8")

    pending = build_input_package(2, config)
    assert pending["status"] == "awaiting_user_inputs"
    assert pending["opportunity_type"] == "speculative_approach"
    assert set(pending["missing_prompts"]) == {"cv_tailoring"}
    assert "master_cv" in pending["missing_inputs"]

    cv_prompt.write_text("Tailor for {company} and {role}: {research_summary}; {job_description}", encoding="utf-8")
    master_cv = tmp_path / "master-cv.docx"
    make_doc(master_cv, "Samuel Anyan master CV")
    config["paths"]["master_cv"] = master_cv
    ready = build_input_package(2, config, role="Analyst", job_description="Verified JD")
    assert ready["status"] == "ready_for_attended_handoff"
    assert "Example Energy" in ready["prompts"]["cv_tailoring"]
    assert "Verified JD" in ready["prompts"]["cv_tailoring"]
    assert ready["sources"][0]["url"] == "https://example.test/source"
    assert expected_filename("Example Energy!", "Cost Engineer", "CV") == (
        "Samuel Anyan_Example Energy_Cost Engineer_CV.docx"
    )
    assert expected_filename("Example Energy", "Cost Engineer", "Cover Letter") == (
        "Samuel Anyan_Example Energy_Cost Engineer_Cover Letter.docx"
    )


def test_docx_validation_checks_company_and_role(tmp_path):
    path = tmp_path / "sample.docx"
    make_doc(path, "Example Energy Analyst")
    assert validate_docx(path, "Example Energy", "Analyst")["bytes"] > 0
    with pytest.raises(DocumentValidationError):
        validate_docx(path, "Different Company")
    with pytest.raises(DocumentValidationError):
        validate_docx(path, "Example Energy", "Director")
    generated = save_text_as_docx("Company-specific text exactly as returned.", tmp_path / "text.docx")
    assert validate_docx(generated, "Company-specific")["bytes"] > 0


def test_supervised_handoff_opens_ui_and_auto_files_new_validated_downloads(tmp_path):
    workbook_path = tmp_path / "dashboard.xlsx"
    make_dashboard(workbook_path, approvals=((2, "Example Energy", "Approved", "Start"),))
    config = make_config(tmp_path, workbook_path)
    write_research(config)
    config["paths"]["master_cv"] = tmp_path / "master-cv.docx"
    make_doc(config["paths"]["master_cv"], "Samuel Anyan master CV")
    cv_prompt = config["paths"]["prompts_dir"] / "cv_tailoring_prompt.md"
    cv_prompt.write_text("Tailor for {company}; research {research_summary}", encoding="utf-8")
    downloads = tmp_path / "Downloads"
    downloads.mkdir()
    old_doc = downloads / "old.docx"
    make_doc(old_doc, "Example Energy old CV Experience Education Skills")
    config["paths"]["downloads_dir"] = downloads
    config["paths"]["database"] = tmp_path / "state.sqlite"
    state = StateStore(config["paths"]["database"])
    state.initialize()
    state.transition(2, "Example Energy", "RESEARCH_READY")
    state.transition(2, "Example Energy", "PENDING_APPROVAL")
    state.transition(2, "Example Energy", "APPROVED")
    state.transition(2, "Example Energy", "APPLICATION_PACKAGE_READY")
    state.transition(2, "Example Energy", "AWAITING_AI_OUTPUT")
    package = build_input_package(2, config, role="Speculative Approach")
    opened_urls = []
    opened_files = []
    service = SupervisedChatGPTHandoff(
        config, state, browser_open=opened_urls.append, open_file=opened_files.append, stable_polls=2
    )
    try:
        session = service.start_for_package(package)
        assert session and session["company_id"] == "2"
        assert opened_urls[0] == "https://chatgpt.com/"
        assert "127.0.0.1" in opened_urls[1]
        page = service._render_page(state.active_handoff())
        assert "Copy CV prompt" in page and "Copy proposal prompt" in page
        assert "Open master CV" in page and "does not submit messages" in page
        service.open_file(config["paths"]["master_cv"])
        assert opened_files == [config["paths"]["master_cv"]]

        cv_download = downloads / "Document.docx"
        letter_download = downloads / "Document (1).docx"
        make_doc(cv_download, "Example Energy Professional Experience Education Technical Skills")
        assert service.monitor_downloads()["completed"] is False
        assert service.monitor_downloads()["completed"] is False
        make_doc(letter_download, "Dear Hiring Team, Example Energy. I am writing about the opportunity. Kind regards.")
        assert service.monitor_downloads()["completed"] is False
        assert service.monitor_downloads()["completed"] is True

        output = config["paths"]["generated_dir"] / "2_Example Energy"
        filed = {path.name for path in output.glob("*.docx")}
        assert "Samuel Anyan_Example Energy_Speculative Approach_CV.docx" in filed
        assert "Samuel Anyan_Example Energy_Speculative Approach_Cover Letter.docx" in filed
        assert old_doc.exists()
        assert not cv_download.exists() and not letter_download.exists()
        assert state.get_handoff(2)["status"] == "completed"
        assert state.get_company(2)["state"] == "DOCUMENTS_READY"
        workbook = load_workbook(workbook_path)
        assert workbook["03 Execution"]["F2"].value == "Documents Ready"
        workbook.close()
    finally:
        service.close()


def test_download_classifier_fails_closed_for_ambiguous_document():
    assert classify_document(Path("document.docx"), "Example Energy consultant information.") is None
    assert snapshot_downloads(Path("/path/that/does/not/exist")) == {}


def test_apollo_csv_import_filters_identity_and_ranks_verified_contacts(tmp_path):
    registry = tmp_path / "registry.csv"
    registry.write_text("id,company,town,sector\n2,Example Energy,London,Energy\n", encoding="utf-8")
    source = tmp_path / "contacts.csv"
    source.write_text(
        "First Name,Last Name,Title,Email,Email Status,Company,LinkedIn URL\n"
        "Ava,Smith,Hiring Manager,ava@example.test,Verified,Example Energy,https://linkedin.com/in/ava\n"
        "No,Match,Director,no@example.test,Verified,Other Company,\n"
        "Unknown,Email,Analyst,,Unknown,Example Energy,\n",
        encoding="utf-8",
    )

    contacts = import_apollo_csv(source, registry, ["Hiring Manager"])

    assert len(contacts) == 2
    assert contacts[0].full_name == "Ava Smith"
    assert contacts[0].email_status == "verified"
    assert contacts[1].contact_status == "no_email"
    assert rank_contact(Contact("2", "Example Energy", "A", role="Hiring Manager"), ["Hiring Manager"]) > 0


def test_graph_provider_uses_only_mock_http_transport(tmp_path):
    calls = []

    def respond(request):
        calls.append((request.method, request.url.path, request.headers.get("authorization")))
        if request.url.path.endswith("/messages"):
            return httpx.Response(201, json={"id": "draft-1"})
        return httpx.Response(202)

    client = httpx.Client(transport=httpx.MockTransport(respond))
    provider = GraphOutlookProvider(lambda: "mock-token", client=client, graph_url="https://graph.mock/v1.0")
    attachment = tmp_path / "cv.docx"
    attachment.write_bytes(b"not a real document; transport fixture only")
    draft_id = provider.create_draft(
        recipient="person@example.test", subject="Preview", body="Mock body", attachments=[attachment]
    )
    provider.send_draft(draft_id)

    assert draft_id == "draft-1"
    assert [item[1] for item in calls] == [
        "/v1.0/me/messages",
        "/v1.0/me/messages/draft-1/attachments",
        "/v1.0/me/messages/draft-1/send",
    ]
    assert all(item[2] == "Bearer mock-token" for item in calls)


def test_gmail_provider_uses_only_mock_http_transport(tmp_path):
    calls = []

    def respond(request):
        calls.append((request.method, request.url.path, request.headers.get("authorization")))
        return httpx.Response(200, json={"id": "gmail-draft"})

    client = httpx.Client(transport=httpx.MockTransport(respond))
    provider = GmailProvider(lambda: "mock-google-token", client=client, api_url="https://gmail.mock/v1")
    attachment = tmp_path / "cv.docx"
    attachment.write_bytes(b"mock attachment")
    draft_id = provider.create_draft(
        recipient="person@example.test", subject="Preview", body="Mock body", attachments=[attachment]
    )
    provider.send_draft(draft_id)

    assert draft_id == "gmail-draft"
    assert [item[1] for item in calls] == [
        "/v1/users/me/drafts", "/v1/users/me/drafts/send"
    ]
    assert all(item[2] == "Bearer mock-google-token" for item in calls)


def test_outreach_batch_requires_verified_contact_approval_and_explicit_send(tmp_path):
    config, state = prepare_outreach(tmp_path)
    contact = state.list_contacts(2)[0]
    assert contact["email_status"] == "verified"
    batch = create_email_batch(config, state)
    provider = MockMailProvider()
    assert len(batch["payload"]["messages"]) == 1
    assert draft_email_batch(config, state, batch["batch_id"], provider) == {
        "drafted": 1, "skipped": 0, "failed": 0
    }
    assert approve_email_batch(state, batch["batch_id"])
    with pytest.raises(BatchSafetyError):
        send_approved_batch(config, state, batch["batch_id"], provider, send_flag=True)
    assert provider.sent == []

    config["mail"]["send_enabled"] = True
    config["mail"]["min_send_pause_seconds"] = 0
    config["mail"]["max_send_pause_seconds"] = 0
    result = send_approved_batch(
        config, state, batch["batch_id"], provider, send_flag=True, sleep=lambda _: None,
        random_uniform=lambda *_: 0,
    )
    assert result == {"sent": 1, "skipped": 0, "failed": 0}
    assert len(provider.sent) == 1
    assert state.daily_sent_count() == 1
    workbook = load_workbook(config["paths"]["workbook"])
    assert workbook["03 Execution"]["D2"].value == "emailed"
    assert workbook["03 Execution"]["E2"].value == "done"
    assert workbook["04 Connections"]["G2"].value == "emailed"
    workbook.close()


def test_outreach_requires_every_recipients_draft_before_batch_approval(tmp_path):
    config, state = prepare_outreach(tmp_path)
    batch = create_email_batch(config, state)
    assert not approve_email_batch(state, batch["batch_id"])


def test_workbook_contact_status_change_after_approval_blocks_send(tmp_path):
    config, state = prepare_outreach(tmp_path)
    batch = create_email_batch(config, state)
    provider = MockMailProvider()
    draft_email_batch(config, state, batch["batch_id"], provider)
    assert approve_email_batch(state, batch["batch_id"])
    workbook = load_workbook(config["paths"]["workbook"])
    workbook["04 Connections"]["G2"] = "unverified"
    workbook.save(config["paths"]["workbook"])
    workbook.close()
    config["mail"]["send_enabled"] = True

    result = send_approved_batch(
        config, state, batch["batch_id"], provider, send_flag=True,
        sleep=lambda _: None, random_uniform=lambda *_: 0,
    )

    assert result["sent"] == 0
    assert result["skipped"] == 1
    assert provider.sent == []


def test_cancel_unsent_batch_deletes_mock_draft_and_allows_new_snapshot(tmp_path):
    config, state = prepare_outreach(tmp_path)
    batch = create_email_batch(config, state)
    provider = MockMailProvider()
    draft_email_batch(config, state, batch["batch_id"], provider)
    assert provider.drafts

    cancel_email_batch(state, batch["batch_id"], provider)
    assert provider.drafts == {}
    assert state.get_batch(batch["batch_id"])["status"] == "cancelled"
    copy_path = config["paths"]["generated_dir"] / "2_Example Energy" / "outreach.json"
    copy = json.loads(copy_path.read_text())
    copy["messages"][0]["body"] = "Revised introduction requiring fresh approval."
    copy_path.write_text(json.dumps(copy))
    replacement = create_email_batch(config, state)

    assert replacement["batch_id"] != batch["batch_id"]
    assert replacement["payload"]["messages"][0]["body"] != batch["payload"]["messages"][0]["body"]


def test_partial_batch_resumes_without_resending_completed_contact(tmp_path):
    config, state = prepare_outreach(tmp_path, contact_count=2)
    config["mail"]["send_enabled"] = True
    config["mail"]["daily_send_limit"] = 1
    config["mail"]["min_send_pause_seconds"] = 0
    config["mail"]["max_send_pause_seconds"] = 0
    batch = create_email_batch(config, state)
    provider = MockMailProvider()
    assert draft_email_batch(config, state, batch["batch_id"], provider)["drafted"] == 2
    assert approve_email_batch(state, batch["batch_id"])

    first = send_approved_batch(
        config, state, batch["batch_id"], provider, send_flag=True,
        sleep=lambda _: None, random_uniform=lambda *_: 0,
    )
    assert first == {"sent": 1, "skipped": 1, "failed": 0}
    assert state.get_batch(batch["batch_id"])["status"] == "partial"
    assert load_workbook(config["paths"]["workbook"])["03 Execution"]["D2"].value == "Approved"

    config["mail"]["daily_send_limit"] = 2
    second = send_approved_batch(
        config, state, batch["batch_id"], provider, send_flag=True,
        sleep=lambda _: None, random_uniform=lambda *_: 0,
    )
    assert second == {"sent": 1, "skipped": 1, "failed": 0}
    assert len(provider.sent) == 2
    assert state.get_batch(batch["batch_id"])["status"] == "completed"
    assert state.daily_sent_count() == 2


def test_stop_file_halts_batch_after_current_send(tmp_path):
    config, state = prepare_outreach(tmp_path, contact_count=2)
    config["mail"]["send_enabled"] = True
    config["mail"]["min_send_pause_seconds"] = 0
    config["mail"]["max_send_pause_seconds"] = 0
    batch = create_email_batch(config, state)

    class StopAfterFirstProvider(MockMailProvider):
        def send_draft(self, provider_id):
            result = super().send_draft(provider_id)
            config["paths"]["stop_file"].write_text("stop", encoding="ascii")
            return result

    provider = StopAfterFirstProvider()
    draft_email_batch(config, state, batch["batch_id"], provider)
    assert approve_email_batch(state, batch["batch_id"])
    result = send_approved_batch(
        config, state, batch["batch_id"], provider, send_flag=True,
        sleep=lambda _: None, random_uniform=lambda *_: 0,
    )
    assert result["sent"] == 1
    assert result["skipped"] == 1
    assert len(provider.sent) == 1


def test_orchestrator_dry_run_does_not_create_state_or_output(tmp_path):
    workbook = tmp_path / "dashboard.xlsx"
    make_dashboard(workbook, approvals=(
        (2, "Example Energy", "Approved", "Start"),
        (3, "Second Energy", "Approved", "Start"),
    ))
    config = make_config(tmp_path, workbook)
    write_research(config)
    config["execution"]["max_companies_per_day"] = 1
    state_path = config["paths"]["database"]

    result = run_approved_queue(config, StateStore(state_path), dry_run=True)

    assert result["planned"] == 1
    assert result["written"] is False
    assert not state_path.exists()
    assert not config["paths"]["generated_dir"].exists()


def test_instance_lock_and_stop_file_are_scoped(tmp_path):
    lock_path = tmp_path / "japp.lock"
    stop_path = tmp_path / "STOP"
    lock = SingleInstanceLock(lock_path)
    lock.acquire()
    with pytest.raises(InstanceAlreadyRunning):
        SingleInstanceLock(lock_path).acquire()
    assert request_stop(stop_path, lock_path)
    lock.release()
    assert not lock_path.exists()
    assert stop_path.exists()


def test_engine_startup_and_graceful_shutdown_with_fake_scheduler(tmp_path):
    workbook = tmp_path / "dashboard.xlsx"
    make_dashboard(workbook)
    config = make_config(tmp_path, workbook)
    class FakeScheduler:
        def __init__(self, **_kwargs):
            self.running = False
            self.jobs = []
        def add_job(self, function, trigger, **kwargs):
            self.jobs.append((function, trigger, kwargs))
        def start(self):
            self.running = True
        def shutdown(self, wait=True):
            self.running = False

    scheduler = None
    def factory(**kwargs):
        nonlocal scheduler
        scheduler = FakeScheduler(**kwargs)
        return scheduler
    def stop_wait(_seconds):
        config["paths"]["stop_file"].write_text("stop", encoding="ascii")

    start_engine(config, scheduler_factory=factory, wait=stop_wait)

    assert scheduler is not None and len(scheduler.jobs) == 5
    assert not config["paths"]["lock_file"].exists()
    assert not config["paths"]["stop_file"].exists()
    assert StateStore(config["paths"]["database"]).services()


def test_github_pull_refuses_dirty_research_without_fetching(tmp_path):
    calls = []

    def runner(command, **kwargs):
        calls.append(command)
        return __import__("subprocess").CompletedProcess(
            command, 0, stdout=" M research/companies/0002.json\n", stderr=""
        )

    with pytest.raises(DirtyResearchTree):
        pull_research_records(tmp_path, runner=runner)
    assert len(calls) == 1
    assert calls[0][:3] == ["git", "status", "--porcelain"]


def test_contact_import_deduplicates_csv_file_and_keeps_registry_private(tmp_path):
    workbook = tmp_path / "dashboard.xlsx"
    make_dashboard(workbook)
    config = make_config(tmp_path, workbook)
    source = config["paths"]["apollo_import_dir"] / "apollo.csv"
    source.parent.mkdir(parents=True)
    source.write_text(
        "First Name,Last Name,Title,Email,Email Status,Company,LinkedIn URL\n"
        "Ava,Smith,Hiring Manager,ava@example.test,Verified,Example Energy,https://linkedin.com/in/ava\n",
        encoding="utf-8",
    )
    state = StateStore(config["paths"]["database"])
    state.initialize()

    first = import_apollo_file(
        source, registry_path=config["paths"]["registry"], workbook_path=workbook,
        state=state, target_titles=["Hiring Manager"],
    )
    second = import_apollo_file(
        source, registry_path=config["paths"]["registry"], workbook_path=workbook,
        state=state, target_titles=["Hiring Manager"],
    )

    assert first["added"] == 1
    assert second["already_imported"] == 1
    assert load_workbook(workbook)["04 Connections"]["C2"].value == "Ava Smith"


def test_contact_monitor_continues_after_one_malformed_csv(tmp_path):
    workbook = tmp_path / "dashboard.xlsx"
    make_dashboard(workbook)
    config = make_config(tmp_path, workbook)
    imports = config["paths"]["apollo_import_dir"]
    imports.mkdir(parents=True)
    (imports / "00-bad.csv").write_text("not,the,required,headers\n1,2,3,4\n", encoding="utf-8")
    (imports / "01-good.csv").write_text(
        "First Name,Last Name,Title,Email,Email Status,Company,LinkedIn URL\n"
        "Ava,Smith,Hiring Manager,ava@example.test,Verified,Example Energy,https://linkedin.com/in/ava\n",
        encoding="utf-8",
    )
    state = StateStore(config["paths"]["database"])
    state.initialize()

    result = import_pending_contacts(config, state)

    assert result == {"files": 2, "added": 1, "skipped": 0, "failed": 1}
    assert state.list_contacts(2)[0]["full_name"] == "Ava Smith"

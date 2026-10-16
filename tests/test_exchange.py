"""Mock-only integration coverage for the private repository workflow."""

import json
from pathlib import Path

import httpx
import pytest
from openpyxl import load_workbook

from test_automation import make_config, make_dashboard, make_doc, write_research
from automation.contact_discovery import discover_contacts
from automation.dashboard import dashboard_copy
from automation.delivery import deliver
from automation.document_generation import build_input_package, expected_filename
from automation.exchange import (collect_outputs, company_folder, export_research, package_inputs, write_json)
from automation.outreach import BatchSafetyError, create_email_batch, validate_final_outreach
from automation.private_repository import PrivateRepository, inside
from automation.providers.apollo import ApolloAPIProvider
from automation.providers.mail import MockMailProvider
from automation.state import StateStore


class FakeRepository:
    repository = "example/japp_data"

    def __init__(self, path):
        self.path = path
        path.mkdir()
        self.published = []

    def publish(self, files, message):
        assert all(p.is_relative_to(self.path) and p.is_file() for p in files)
        self.published.append((files, message))
        return True


@pytest.fixture
def setup_exchange(tmp_path):
    workbook = tmp_path / "dashboard.xlsx"
    make_dashboard(workbook)
    config = make_config(tmp_path, workbook)
    config["paths"]["master_cv"] = tmp_path / "master.docx"
    make_doc(config["paths"]["master_cv"], "Samuel Anyan master CV facts")
    (config["paths"]["prompts_dir"] / "cv_tailoring_prompt.md").write_text("Tailor for {company}: {research_summary}")
    repo = FakeRepository(tmp_path / "private")
    research = write_research(config)
    raw = json.loads(research.read_text())
    raw["website"] = "https://example.test"
    write_json(repo.path / "research/companies/2.json", raw)
    state = StateStore(config["paths"]["database"])
    state.initialize()
    return config, repo, state


def finish_outputs(config, repo, state):
    package_inputs(config, repo, state)
    folder = company_folder(repo.path, "2", "Example Energy")
    state.save_contact("2", "Example Energy", full_name="A Person", role="Hiring Manager",
                       email="person@example.test", email_status="verified")
    from automation.workbook import write_contacts
    write_contacts(config["paths"]["workbook"], "2", [{"full_name": "A Person", "role": "Hiring Manager",
                    "email": "person@example.test", "email_status": "verified", "contact_status": "found"}])
    make_doc(folder / "outputs/random.docx", "Samuel Anyan Professional Experience Education Skills")
    make_doc(folder / "outputs/download (2).docx", "Dear Example Energy, I am writing. Kind regards.")
    write_json(folder / "outputs/outreach.json", {"messages": [{"recipient": "person@example.test",
               "subject": "An introduction to Example Energy", "body": "Hi A,\nI would welcome a conversation.\nSamuel"}]})
    package = json.loads((folder / "inputs/input_package.json").read_text())
    manifest = {"schema_version": "1.0", "ready": True, "company_id": "2", "company": "Example Energy",
                "role": package["role"], "input_hash": package["input_hash"], "cv": "random.docx",
                "cover_letter": "download (2).docx", "outreach": "outreach.json"}
    write_json(folder / "outputs/manifest.json", manifest)
    return folder, manifest


def test_research_export_is_stable_and_updates_only_pipeline_status(setup_exchange):
    config, repo, _ = setup_exchange
    result = export_research(config, repo)
    assert result["companies"] == 1
    assert result["batch_id"] == export_research(config, repo)["batch_id"]
    book = load_workbook(config["paths"]["workbook"])
    assert book["03 Execution"]["D2"].value == "Approved"
    assert book["03 Execution"]["E2"].value == "Start"
    assert book["03 Execution"]["F2"].value == "Exported"
    book.close()
    request = json.loads((repo.path / f"requests/{result['batch_id']}/companies.json").read_text())
    assert request["companies"][0]["company_id"] == "2"


def test_research_scope_excludes_previously_selected_companies(setup_exchange):
    config, repo, _ = setup_exchange
    assert export_research(config, repo, company_ids={"16", "19"}) == {"companies": 0, "published": False}
    assert repo.published == []
    result = export_research(config, repo, company_ids={"2"})
    request = json.loads((repo.path / f"requests/{result['batch_id']}/companies.json").read_text())
    assert [r['company_id'] for r in request['companies']] == ['2']


def test_input_package_is_portable_and_tracks_original_version(setup_exchange):
    config, repo, state = setup_exchange
    result = package_inputs(config, repo, state)
    assert result["packaged"] == 1
    folder = company_folder(repo.path, "2", "Example Energy")
    package = json.loads((folder / "inputs/input_package.json").read_text())
    assert package["master_cv_path"] == "master_cv.docx"
    assert "package_path" not in package and "output_dir" not in package
    assert state.published_input_hash("2") == package["input_hash"]
    assert state.get_company("2")["state"] == "AWAITING_AI_OUTPUT"
    book = load_workbook(config["paths"]["workbook"])
    assert book["03 Execution"]["F2"].value == "Researched"
    book.close()
    assert (folder / "inputs/master_cv.docx").read_bytes() == config["paths"]["master_cv"].read_bytes()
    assert "speculative" in (folder / "inputs/TASK.md").read_text()
    assert package_inputs(config, repo, state)["packaged"] == 1
    assert state.published_input_hash("2") == package["input_hash"]


def test_package_scope_leaves_unselected_company_untouched(setup_exchange):
    config, repo, state = setup_exchange
    assert package_inputs(config, repo, state, company_ids={"4", "5"})["packaged"] == 0
    assert repo.published == []
    assert state.get_company("2") is None
    assert not (repo.path / "companies").exists()
    assert package_inputs(config, repo, state, company_ids={"2"})["packaged"] == 1


def test_package_removes_legacy_prompts_and_uses_only_modular_spec(setup_exchange):
    config, repo, state = setup_exchange
    folder = company_folder(repo.path, "2", "Example Energy") / "inputs"
    folder.mkdir(parents=True)
    for name in ("proposal_prompt.md", "outreach_template.md"):
        (folder / name).write_text("Obsolete writing instructions")
    # FakeRepository expects published files to exist; verify deletion staging separately.
    published = []
    repo.publish = lambda files, message: published.extend(files)
    assert package_inputs(config, repo, state)["packaged"] == 1
    assert {p.name for p in folder.glob("*prompt*.md")} == {"cv_tailoring_prompt.md"}
    assert not (folder / "outreach_template.md").exists()
    assert (folder / "proposal_prompt.md") in published
    package = json.loads((folder / "input_package.json").read_text())
    assert set(package["prompts"]) == {"cv_tailoring"}


def test_modular_prompt_placeholder_is_detected(setup_exchange):
    config, repo, state = setup_exchange
    (config["paths"]["prompts_dir"] / "cv_tailoring_prompt.md").write_text(
        "# Proposal prompt\n<!-- Add your own proposal/introduction prompt here. -->")
    assert "cv_tailoring" in build_input_package("2", config)["missing_prompts"]
    assert package_inputs(config, repo, state)["packaged"] == 0
    assert not repo.published


def test_collect_renames_pair_updates_states_and_is_idempotent(setup_exchange):
    config, repo, state = setup_exchange
    finish_outputs(config, repo, state)
    assert collect_outputs(config, repo, state)["collected"] == ["2"]
    assert state.get_company("2")["state"] == "CONTACTS_READY"
    assert collect_outputs(config, repo, state)["collected"] == ["2"]
    folder = config["paths"]["generated_dir"] / "2_Example Energy"
    assert (folder / expected_filename("Example Energy", "Speculative Approach", "CV")).is_file()
    book = load_workbook(config["paths"]["workbook"])
    assert book["03 Execution"]["F2"].value == "Documents Ready"
    assert book["03 Execution"]["D2"].value == "Approved"
    assert book["03 Execution"]["E2"].value == "Start"
    book.close()


@pytest.mark.parametrize("change", ["hash", "identity", "path", "role", "tampered_input"])
def test_collect_rejects_stale_or_mismatched_outputs(setup_exchange, change):
    config, repo, state = setup_exchange
    folder, manifest = finish_outputs(config, repo, state)
    if change == "hash": manifest["input_hash"] = "old"
    elif change == "identity": manifest["company"] = "Other Company"
    elif change == "path": manifest["cv"] = "../inputs/master_cv.docx"
    elif change == "role": manifest["role"] = "Wrong Role"
    else:
        p = json.loads((folder / "inputs/input_package.json").read_text())
        p["input_hash"] = manifest["input_hash"] = "tampered"
        write_json(folder / "inputs/input_package.json", p)
    write_json(folder / "outputs/manifest.json", manifest)
    with pytest.raises(ValueError): collect_outputs(config, repo, state)
    assert state.get_company("2")["state"] == "AWAITING_AI_OUTPUT"


def test_incomplete_manifest_and_revoked_approval_do_not_advance(setup_exchange):
    config, repo, state = setup_exchange
    folder, manifest = finish_outputs(config, repo, state)
    manifest["ready"] = False
    write_json(folder / "outputs/manifest.json", manifest)
    assert not collect_outputs(config, repo, state)["collected"]
    manifest["ready"] = True
    write_json(folder / "outputs/manifest.json", manifest)
    book = load_workbook(config["paths"]["workbook"])
    book["03 Execution"]["D2"] = "Rejected"
    book.save(config["paths"]["workbook"])
    book.close()
    assert not collect_outputs(config, repo, state)["collected"]


def test_delivery_uses_final_copy_not_prompt_and_requires_send_flag(setup_exchange):
    config, repo, state = setup_exchange
    finish_outputs(config, repo, state)
    provider = MockMailProvider()
    config["mail"]["min_send_pause_seconds"] = config["mail"]["max_send_pause_seconds"] = 0
    result = deliver(config, repo, state, provider=provider)
    assert result["drafts"]["drafted"] == 1 and result["sent"] == 0
    book = load_workbook(config["paths"]["workbook"])
    assert book["03 Execution"]["F2"].value == "Email Drafted"
    book.close()
    batch = state.get_batch(result["batch_id"])
    assert batch["payload"]["messages"][0]["subject"] == "An introduction to Example Energy"
    assert "{recipient_name}" not in batch["payload"]["messages"][0]["body"]
    with pytest.raises(ValueError): deliver(config, repo, state, send=True, provider=provider)
    config["mail"]["send_enabled"] = True
    sent = deliver(config, repo, state, send=True, provider=provider)
    assert sent["sent"] == 1
    assert state.get_company("2")["state"] == "COMPLETED"
    book = load_workbook(config["paths"]["workbook"])
    assert book["03 Execution"]["F2"].value == "Completed"
    book.close()
    assert deliver(config, repo, state, send=True, provider=provider)["batch_id"] is None


class FakeApollo:
    def __init__(self, fail=False):
        self.calls = 0
        self.fail = fail

    def search(self, company_domain, titles, limit):
        return [{"id": "p1", "organization": {"name": "Example Energy"}}]

    def enrich(self, person_id):
        self.calls += 1
        if self.fail: raise RuntimeError("uncertain transport")
        return {"person": {"id": person_id, "name": "A Person", "title": "Hiring Manager",
                "email": "person@example.test", "email_status": "verified",
                "organization": {"primary_domain": "example.test"}}}


def test_apollo_reveals_are_cached_and_preserve_suppression(setup_exchange):
    config, repo, state = setup_exchange
    package_inputs(config, repo, state)
    api = FakeApollo()
    kwargs = dict(reveal_emails=True, max_reveals=1, provider=api)
    assert discover_contacts(config, repo, state, **kwargs)["contacts_saved"] == 1
    book = load_workbook(config["paths"]["workbook"])
    assert book["03 Execution"]["F2"].value == "Contacts Ready"
    book.close()
    state.set_contact_suppressed(state.list_contacts("2")[0]["contact_id"])
    assert discover_contacts(config, repo, state, **kwargs)["reveals"] == 0
    assert api.calls == 1 and state.list_contacts("2")[0]["suppressed"]
    saved = json.loads((company_folder(repo.path, "2", "Example Energy") / "contacts.json").read_text())
    assert saved["company_id"] == "2" and saved["contacts"][0]["suppressed"]


def test_apollo_uncertain_request_is_not_charged_again(setup_exchange):
    config, repo, state = setup_exchange
    package_inputs(config, repo, state)
    api = FakeApollo(fail=True)
    kwargs = dict(reveal_emails=True, max_reveals=1, provider=api)
    with pytest.raises(RuntimeError): discover_contacts(config, repo, state, **kwargs)
    assert discover_contacts(config, repo, state, **kwargs)["reveals"] == 0
    assert api.calls == 1


def test_supported_apollo_requests_only_work_email(setup_exchange, monkeypatch):
    config, _, _ = setup_exchange
    monkeypatch.setenv("APOLLO_API_KEY", "mock-only-key")
    calls = []
    def handler(request):
        calls.append((request.url.path, json.loads(request.content)))
        return httpx.Response(200, json={"people": [], "person": {}})
    provider = ApolloAPIProvider(config, client=httpx.Client(transport=httpx.MockTransport(handler)))
    provider.search("example.test", ["Hiring Manager"], 5)
    provider.enrich("p1")
    assert calls[0][0] == "/api/v1/mixed_people/api_search"
    assert calls[1][1] == {"id": "p1", "reveal_personal_emails": False, "reveal_phone_number": False}


def test_private_repository_rejects_public_remote_before_writing(setup_exchange, monkeypatch):
    config, _, _ = setup_exchange
    config["paths"]["private_repo_dir"] = config["paths"]["generated_dir"] / "must-not-exist"
    import automation.private_repository as module
    monkeypatch.setattr(module, "command", lambda *a, **kw: json.dumps({"nameWithOwner": "sa-anyan/japp_data", "isPrivate": False}))
    with pytest.raises(ValueError): PrivateRepository(config).sync()
    assert not config["paths"]["private_repo_dir"].exists()


def test_exchange_paths_reject_symlinks_and_traversal(tmp_path):
    (tmp_path / "link").symlink_to(tmp_path)
    for name in ("../private.txt", "/absolute.txt", ".git/config", "link/file.docx"):
        with pytest.raises(ValueError): inside(tmp_path, name)


def test_dashboard_start_control_is_in_new_copy_only(setup_exchange, tmp_path):
    config, _, _ = setup_exchange
    source = config["paths"]["workbook"]
    before = source.read_bytes()
    output = dashboard_copy(source, tmp_path / "dashboard-with-start.xlsx")
    assert source.read_bytes() == before
    book = load_workbook(output)
    controls = [v for v in book["03 Execution"].data_validations.dataValidation if "E2" in v]
    assert len(controls) == 1 and 'Start,Done,Not required' in controls[0].formula1
    assert controls[0].showDropDown is False
    book.close()
    with pytest.raises(ValueError): dashboard_copy(source, source)
    with pytest.raises(ValueError): dashboard_copy(source, output)


def test_apollo_budget_is_global_and_old_employer_match_is_skipped(setup_exchange):
    config, repo, state = setup_exchange
    package_inputs(config, repo, state)
    api = FakeApollo()
    api.search = lambda *args: [
        {"id": "wrong", "organization": {"name": "Previous Employer"}},
        {"id": "p1", "organization": {"name": "Example Energy"}},
        {"id": "p2", "organization": {"name": "Example Energy"}},
    ]
    result = discover_contacts(config, repo, state, reveal_emails=True, max_reveals=1, provider=api)
    assert result["reveals"] == 1 and api.calls == 1
    assert any("budget" in item for item in result["waiting"])
    assert state.enrichment("wrong") is None and state.enrichment("p2") is None


def test_collect_will_not_overwrite_an_existing_local_cv(setup_exchange):
    config, repo, state = setup_exchange
    finish_outputs(config, repo, state)
    local = config["paths"]["generated_dir"] / "2_Example Energy"
    original = local / expected_filename("Example Energy", "Speculative Approach", "CV")
    make_doc(original, "Original user edited CV")
    before = original.read_bytes()
    with pytest.raises(ValueError): collect_outputs(config, repo, state)
    assert original.read_bytes() == before
    assert not (local / "outreach.json").exists()


def test_changed_inputs_after_downstream_work_require_explicit_reset(setup_exchange):
    config, repo, state = setup_exchange
    finish_outputs(config, repo, state)
    collect_outputs(config, repo, state)
    (config["paths"]["prompts_dir"] / "cv_tailoring_prompt.md").write_text("A changed prompt")
    with pytest.raises(ValueError, match="downstream"):
        package_inputs(config, repo, state)


def test_final_email_rejects_unverified_recipient_and_placeholder():
    message = {"recipient": "person@example.test", "subject": "Hello", "body": "Hi [Recipient First Name]"}
    contacts = [{"email": "person@example.test", "email_status": "verified", "suppressed": False}]
    with pytest.raises(BatchSafetyError): validate_final_outreach({"messages": [message]}, contacts)
    message["body"] = "A final introduction"
    contacts[0]["email_status"] = "unknown"
    with pytest.raises(BatchSafetyError): validate_final_outreach({"messages": [message]}, contacts)


def test_dashboard_cli_activates_only_new_copy(setup_exchange, monkeypatch, tmp_path):
    config, _, _ = setup_exchange
    from automation import cli
    config["config_path"] = tmp_path / "config.json"
    config["config_path"].write_text(json.dumps({"mail": {"send_enabled": False}}))
    monkeypatch.setattr(cli, "load_config", lambda *args: config)
    destination = tmp_path / "active.xlsx"
    original = config["paths"]["workbook"].read_bytes()
    assert cli.main(["dashboard-copy", "--output", str(destination), "--activate"]) == 0
    saved = json.loads(config["config_path"].read_text())
    assert saved["paths"]["workbook"] == str(destination)
    assert saved["mail"]["send_enabled"] is False
    assert config["paths"]["workbook"].read_bytes() == original


def test_cv_identity_is_required_but_target_company_is_not(tmp_path):
    from automation.document_generation import validate_docx, DocumentValidationError
    cv = tmp_path / 'cv.docx'
    make_doc(cv, 'Samuel Anyan Professional Experience Education Skills')
    assert validate_docx(cv, 'Target Company', document_kind='cv')
    with pytest.raises(DocumentValidationError, match='company'):
        validate_docx(cv, 'Target Company')
    make_doc(cv, 'Another Person Professional Experience Education Skills')
    with pytest.raises(DocumentValidationError, match='identity'):
        validate_docx(cv, 'Target Company', document_kind='cv')

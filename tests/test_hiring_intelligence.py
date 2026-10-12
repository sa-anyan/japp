import csv
import json
import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook

from japp_hiring_intelligence import (
    SHEET, canonical, classify, generate_queries, import_records,
    load_config, save_leads,
)

CONFIG = {
    "hiring_signals": ["#hiring", "We're hiring", "Join my team"],
    "keywords": ["Energy Analyst", "Data Analyst", "Python", "Market Risk"],
    "seniority": ["Graduate", "Junior", "Analyst"],
    "location": "UK",
    "posted_within_days": 7,
    "minimum_score": 45,
    "max_queries": 24,
}
TODAY = date(2026, 10, 13)


class HiringIntelligenceTests(unittest.TestCase):
    def test_queries_have_signal_keyword_and_role_seniority(self):
        queries = generate_queries(CONFIG)
        self.assertGreater(len(queries), 8)
        self.assertEqual(len(queries), len(set(queries)))
        for q in queries:
            self.assertIn("site:linkedin.com/posts/", q)
            self.assertTrue(any(s in q for s in CONFIG["hiring_signals"]))
            self.assertTrue(any(s in q for s in CONFIG["keywords"]))
            self.assertTrue(any(s in q for s in CONFIG["seniority"]))

    def test_director_posting_junior_role_is_not_excluded(self):
        lead = classify({
            "post_url": "https://www.linkedin.com/posts/example-123?utm_source=feed",
            "title": "Junior Energy Analyst",
            "snippet": "#hiring Join my team: Junior Energy Analyst London",
            "poster_role": "Senior Director of Trading",
            "poster": "Example Person",
            "post_date": "2026-10-12",
            "location": "London, UK",
        }, CONFIG, today=TODAY)
        self.assertEqual(lead["Review Status"], "Review")
        self.assertEqual(lead["Target Seniority"], "Junior")
        self.assertEqual(lead["Contact Type"], "Possible hiring manager")
        self.assertEqual(lead["Contact Confidence"], "Unverified")
        self.assertEqual(lead["Sponsorship"], "Unknown")
        self.assertNotIn("utm_source", lead["Post URL"])

    def test_senior_vacancy_and_no_sponsorship_are_excluded(self):
        base = {
            "post_url": "https://www.linkedin.com/posts/senior-role",
            "title": "Senior Data Analyst",
            "snippet": "#hiring Senior Data Analyst in London",
        }
        self.assertEqual(classify(base, CONFIG, today=TODAY)["Review Status"], "Exclude")
        base["title"] = "Junior Data Analyst"
        base["snippet"] = "#hiring Junior Data Analyst - no visa sponsorship"
        lead = classify(base, CONFIG, today=TODAY)
        self.assertEqual(lead["Sponsorship"], "No")
        self.assertEqual(lead["Review Status"], "Exclude")

    def test_date_freshness_and_missing_date_are_not_confused(self):
        base = {
            "post_url": "https://www.linkedin.com/posts/valid",
            "title": "Graduate Data Analyst",
            "snippet": "#hiring Graduate Data Analyst",
        }
        self.assertEqual(classify(base, CONFIG, today=TODAY)["Review Status"], "Review")
        old = dict(base, post_date="2026-09-01")
        self.assertEqual(classify(old, CONFIG, today=TODAY)["Review Status"], "Exclude")
        self.assertIn("post date unverified", classify(base, CONFIG, today=TODAY)["Why Matched"])

    def test_excel_preserves_other_sheets_and_manual_outreach_edits(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "tracker.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.title = "Applications"
            sheet["A1"] = "My existing application data"
            book.save(path)
            row = {
                "post_url": "https://www.linkedin.com/posts/grad-1?tracking=1",
                "title": "Graduate Energy Analyst",
                "snippet": "#hiring Graduate Energy Analyst - UK",
                "company": "Example Energy",
                "post_date": "2026-10-12",
            }
            one = classify(row, CONFIG, today=TODAY)
            result = save_leads(path, [one])
            self.assertEqual(result["added"], 1)
            book = load_workbook(path)
            self.assertEqual(book["Applications"]["A1"].value, "My existing application data")
            hiring = book[SHEET]
            header = {c.value: c.column for c in hiring[1]}
            hiring.cell(2, header["Outreach Status"]).value = "Messaged"
            hiring.cell(2, header["Notes"]).value = "Follow up next week"
            book.save(path)
            # A second search should deduplicate by normalised URL.
            updated = classify(dict(row, post_url=row["post_url"].split("?")[0]),
                               CONFIG, today=TODAY)
            result = save_leads(path, [updated])
            self.assertEqual(result["updated"], 1)
            book = load_workbook(path)
            hiring = book[SHEET]
            self.assertEqual(hiring.max_row, 2)
            self.assertEqual(hiring.cell(2, header["Outreach Status"]).value, "Messaged")
            self.assertEqual(hiring.cell(2, header["Notes"]).value, "Follow up next week")
            self.assertIsNotNone(hiring.cell(2, header["Post URL"]).hyperlink)

    def test_import_csv_and_json(self):
        with tempfile.TemporaryDirectory() as td:
            csv_file = Path(td) / "posts.csv"
            with csv_file.open("w", newline="", encoding="utf-8") as f:
                writer = csv.DictWriter(f, fieldnames=["post_url", "title", "snippet"])
                writer.writeheader()
                writer.writerow({"post_url": "https://linkedin.com/posts/1",
                                 "title": "Junior Market Risk",
                                 "snippet": "#hiring Junior Market Risk"})
            self.assertEqual(len(import_records(csv_file)), 1)
            json_file = Path(td) / "posts.json"
            json_file.write_text(json.dumps({"posts": [{"post_url": "https://linkedin.com/posts/2"}]}))
            self.assertEqual(len(import_records(json_file)), 1)

    def test_reject_non_http_post_urls(self):
        self.assertEqual(canonical("file:///Users/sam/private"), "")
        self.assertEqual(canonical("javascript:alert(1)"), "")
        self.assertEqual(canonical("https://www.linkedin.com/posts/abc?utm=1"),
                         "https://linkedin.com/posts/abc")


if __name__ == "__main__":
    unittest.main()
